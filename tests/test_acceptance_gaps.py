"""验收矩阵（A01–A20）的**缺口补测**：逐项对应文档第 4.1 节。

已有的覆盖在哪（本文件不重复）：

* A01/A02/A03/A04/A05/A06/A07 —— ``tests/test_wps_cloud.py`` 与
  ``tests/test_wps_cloud_safety.py``（增量累加、重复上传零写入、协作者 0 保护、
  多槽位、批次日期闸门、排序前回滚、排序后不误删）；
* A11/A17（部分）—— ``tests/test_wps_recovery.py``；
* A13/A14/A15（阻断侧）/A16/A18 —— ``tests/test_sss_journal.py``。

本文件补的是审计确认的缺口 + 本轮修复的回归锁：

* A08 预览后本地文件变化 → 上传前拒绝（重算 SHA-256）
* A09 预览令牌过期 / 关闭后重新开启的复用
* A10 意图日志损坏 → 结构化拒绝且零云端写入
* A12 闪时送日志不可写 → 端到端零 POST
* A13b ``inflight`` 残留（进程被杀）→ 下次运行阻断
* A15b 换账号 / 换平台的**处置**必须被拒绝（快照作用域绑定）
* A16b ``scan_failed`` 不能支撑 ``station_absent``，且仍保持阻断
* A17b 核对快照过期 → 拒绝解除
* A19 浏览器（Playwright）路径：POST 前留痕 + 异常归类与接口路径一致
* 新增：未处置的 ``uncertain`` 必须阻断同日期 + 同表重传（重复累加的最后防线）
* 新增：锁拿不到 / 日志目录不可写 → fail-closed
"""
from __future__ import annotations

import datetime as dt
import json
import sys
import threading
import types
from pathlib import Path

import pytest

from app import sss as sss_module
from app import sss_review as review
from app.sss import run_sss_job
from app.sss_journal import (append_records, batch_key,
                             blocking_state, load_journal)
from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
from app.wps_journal import SyncJournal, journal_path_for
from tests.test_sss_journal import (_FakeClient, _config, _entry, _fetch_queue,
                                    _id_of, _meta, _write_sss_excel)
from tests.test_wps_cloud import BASE_HEADER, FakeCli, make_grid

TARGET = dt.date(2026, 9, 11)


# ======================================================================
# A08：预览之后本地文件变化 → 上传必须在写入前拒绝
# ======================================================================

def test_upload_rejects_when_the_local_file_content_changed(tmp_path, monkeypatch):
    """只比计划指纹不够：任何内容变化都要让预览令牌作废（用户心智与文档要求）。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"original-content")
    bridge._config.excel_path = excel

    calls: list[str] = []
    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: calls.append("apply") or {
                            "sheets": [], "written": 0, "failed": 0})

    preview = bridge.wps_preview()
    assert preview["ok"] is True and preview["preview_id"]

    # 预览之后改文件内容（模拟用户在 Excel 里又编辑过并保存）
    excel.write_bytes(b"changed-content-after-preview")

    got = bridge.wps_upload(preview["preview_id"])

    assert got["ok"] is False
    assert got["code"] == "preview_changed"
    assert "local_file" in (got.get("changed") or [])
    assert got["execution_summary"]["proven_no_write"] is True
    assert calls == [], "文件变了就绝不能调用 apply_plan"


def test_upload_accepts_when_the_file_is_untouched(tmp_path, monkeypatch):
    """反向锁：文件没变时不能因为新加的指纹校验而误拒。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"stable-content")
    bridge._config.excel_path = excel

    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: {"sheets": [], "written": 0, "failed": 0,
                                         "proven_no_write": True})

    preview = bridge.wps_preview()
    got = bridge.wps_upload(preview["preview_id"])
    assert got["ok"] is True


# ======================================================================
# A09：令牌过期 / 关闭后重新开启都拒绝
# ======================================================================

def test_preview_token_expires_after_ttl():
    """直接测令牌表：过期后既不能查成 valid，也不能被消费。"""
    from app.wps_preview import PreviewStore

    store = PreviewStore(ttl_seconds=600)
    record = store.create(local_sha256="sha", context={}, context_fingerprint="ctx",
                          plan=[], plan_fingerprint="plan", summary={}, text="",
                          tables=[], blocked=[], warnings=[], target_date="2026-09-11",
                          target_tables={}, now=1_000.0)

    assert store.get(record.preview_id, now=1_000.0)[1] == ""
    assert store.get(record.preview_id, now=1_599.0)[1] == "", "600 秒内仍有效"
    assert store.get(record.preview_id, now=1_601.0)[1] == "preview_expired"
    assert store.consume(record.preview_id, now=1_601.0)[1] == "preview_expired"


def test_upload_rejects_an_expired_token(tmp_path, monkeypatch):
    """桥接层：过期令牌必须返回 preview_expired 且不调用 apply_plan。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"x")
    bridge._config.excel_path = excel

    calls: list[str] = []
    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: calls.append("apply") or {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    preview = bridge.wps_preview()
    pid = preview["preview_id"]

    # 手动把令牌改成 1 秒 TTL 并推到过期（不改生产 TTL）。
    record, code = bridge._previews.get(pid)
    assert code == "" and record is not None
    record.expires_at = 1.0

    got = bridge.wps_upload(pid)

    assert got["ok"] is False and got["code"] == "preview_expired"
    assert got["execution_summary"]["proven_no_write"] is True
    assert calls == []


def test_reenabling_sync_does_not_revive_old_tokens(tmp_path, monkeypatch):
    """关闭 → 重新开启后，关闭前拿到的 preview_id 依然不能用。

    这条比"关闭窗口内被拒"更强：如果只是靠 `wps_enabled=False` 短路，
    重新开启后旧令牌就会复活并真的写云端。
    """
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"x")
    bridge._config.excel_path = excel

    calls: list[str] = []
    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: calls.append("apply") or {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    preview = bridge.wps_preview()
    pid = preview["preview_id"]

    assert bridge.save_wps_config({"enabled": False})["ok"] is True
    assert bridge.save_wps_config({"enabled": True})["ok"] is True

    got = bridge.wps_upload(pid)

    assert got["ok"] is False, "重新开启后旧令牌不能复活"
    assert got["code"] in ("preview_invalidated", "preview_not_found")
    assert calls == []


# ======================================================================
# A10：意图日志损坏 → 结构化拒绝 + 零云端写入
# ======================================================================

def _broken_journal_bridge(tmp_path, monkeypatch):
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"x")
    bridge._config.excel_path = excel

    calls: list[str] = []
    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: calls.append("apply") or {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    return bridge, calls


def _corrupt_the_journal(bridge) -> Path:
    ledger, journal = bridge._wps_ledger_and_journal()
    path = journal_path_for(ledger.path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("{ 不是合法 JSON", encoding="utf-8")
    return path


def test_corrupt_journal_blocks_preview_without_raising(tmp_path, monkeypatch):
    """预览也必须拒绝：读不出"上次写到哪"就不该继续往下走。"""
    bridge, calls = _broken_journal_bridge(tmp_path, monkeypatch)
    _corrupt_the_journal(bridge)

    got = bridge.wps_preview()          # 不能抛异常

    assert got["ok"] is False
    assert got["code"] == "local_state_blocked"
    assert got["status"] == "blocked"
    assert got["execution_summary"]["proven_no_write"] is True
    assert calls == []


def test_corrupt_journal_blocks_upload_without_raising(tmp_path, monkeypatch):
    """A10 的核心：日志坏了要在写入之前拒绝，且返回结构化结果。"""
    bridge, calls = _broken_journal_bridge(tmp_path, monkeypatch)
    preview = bridge.wps_preview()
    assert preview["ok"] is True
    pid = preview["preview_id"]

    _corrupt_the_journal(bridge)
    got = bridge.wps_upload(pid)        # 不能抛异常

    assert got["ok"] is False
    assert got["code"] == "local_state_blocked"
    assert got["execution_summary"]["proven_no_write"] is True
    assert calls == [], "日志不可读时一行都不能写"


def test_corrupt_ledger_blocks_upload_without_raising(tmp_path, monkeypatch):
    """账本（幂等锚点）损坏同样失败关闭 —— 否则会把同一批餐重复加一遍。"""
    bridge, calls = _broken_journal_bridge(tmp_path, monkeypatch)
    ledger, _journal = bridge._wps_ledger_and_journal()
    ledger.path.parent.mkdir(parents=True, exist_ok=True)
    ledger.path.write_text("{ 坏账本", encoding="utf-8")

    got = bridge.wps_preview()

    assert got["ok"] is False and got["code"] == "local_state_blocked"
    assert calls == []


# ======================================================================
# 新增（本轮修复）：未处置的 uncertain 必须阻断同日期 + 同表重传
# ======================================================================

@pytest.fixture
def wps_env(tmp_path: Path):
    ledger = SyncLedger(tmp_path / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))
    return ledger, journal


def _wps_plan(cli, meals: int = 6):
    orders = [CloudOrder("东湖中餐", "张", "小", "111", "中餐", "经济", meals,
                         row=3, rows=(3,))]
    return build_plan(cli, local_orders={"东湖中餐": orders},
                      tables={"东湖中餐": {"file_id": "F1"}},
                      target=TARGET, ledger=None)


def test_unresolved_wps_write_blocks_the_same_batch_reupload(wps_env):
    """上一次写入结果未知（verify_failed → uncertain）后，再上传同表必须被拒。

    这是重复累加的**最后一道防线**：如果那次其实写成功了，这一次会把同一批餐
    再加一遍。``retire_guarded`` 与 pending 都必须挡住。
    """
    ledger, journal = wps_env
    cli = FakeCli(make_grid(BASE_HEADER, [{0: "张", 2: "111", 7: "3"}]),
                  corrupt_write=True)
    first = apply_plan(cli, _wps_plan(cli), ledger=ledger, marker_enabled=False,
                       journal=journal)
    assert first["sheets"][0]["status"] == "verify_failed"

    cli2 = FakeCli(make_grid(BASE_HEADER, [{0: "张", 2: "111", 7: "3"}]))
    second = apply_plan(cli2, _wps_plan(cli2), ledger=ledger, marker_enabled=False,
                        journal=SyncJournal(journal_path_for(ledger.path)))

    assert second["sheets"][0]["status"] == "stale_batch"
    assert "未处置的同目标写入" in second["sheets"][0]["reason"]
    assert not cli2.writes and not cli2.inserts, "有未处置记录时一个写请求都不该发"
    assert second["proven_no_write"] is True


def test_verified_write_does_not_block_the_next_batch(wps_env):
    """反向锁：正常成功后同一张表仍然可以继续上传（否则功能不可用）。"""
    ledger, journal = wps_env
    cli = FakeCli(make_grid(BASE_HEADER, [{0: "张", 2: "111", 7: "3"}]))
    first = apply_plan(cli, _wps_plan(cli), ledger=ledger, marker_enabled=False,
                       journal=journal)
    assert first["sheets"][0]["status"] == "ok"

    cli2 = FakeCli(make_grid(BASE_HEADER, [
        {0: "张", 2: "111", 4: "1", 5: "中餐", 6: "经济", 7: "9",
         8: "=SUM(D3)", 9: "=H3-I3"}]))
    second = apply_plan(cli2, _wps_plan(cli2), ledger=ledger, marker_enabled=False,
                        journal=SyncJournal(journal_path_for(ledger.path)))
    assert second["sheets"][0]["status"] in ("ok", "noop"), second["sheets"]


def test_blocked_result_reports_proven_no_write_and_zero_counts(wps_env):
    """§2.1.4：整批阻断必须明确报告零写入、零账本更新、可证明未写入。"""
    ledger, _journal = wps_env
    cli = FakeCli(make_grid(BASE_HEADER, []))
    orders = [CloudOrder("东湖中餐", "张", "小", "111", "中餐", "经济", 6, row=3,
                         rows=(3,), weekday_marks=("周五",))]
    plans = build_plan(cli, local_orders={"东湖中餐": orders},
                       tables={"东湖中餐": {"file_id": "F1"}},
                       target=dt.date(2026, 9, 12), ledger=ledger)
    reads_before = len(cli.reads)

    result = apply_plan(cli, plans, ledger=ledger, marker_enabled=False)

    assert result["proven_no_write"] is True
    assert result["cloud_writes"] == 0
    assert result["ledger_updates"] == 0
    sheet = result["sheets"][0]
    assert sheet["status"] == "blocked"
    assert sheet["cloud_writes"] == 0 and sheet["ledger_updates"] == 0
    assert len(cli.reads) == reads_before, "被闸门拒绝的表连只读请求都不该发"


# ======================================================================
# A12：闪时送日志不可写 → 端到端零 POST
# ======================================================================

def test_run_sss_job_sends_no_post_when_the_journal_cannot_be_written(
        tmp_path, monkeypatch):
    """端到端：日志写不下去时，一个 POST 都不许发（而不是"先发再补记"）。"""
    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", sss_max_workers=1)
    client = _FakeClient()
    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: client)

    def boom(*_a, **_k):
        raise OSError("磁盘满")

    monkeypatch.setattr(sss_module, "append_records", boom)
    logs: list[str] = []

    result = run_sss_job(cfg, threading.Event(), progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")

    assert client.posts == [], "留痕失败就绝不能发 POST"
    assert result["created"] == 0
    # 一次都没提交过 → 明确是"本批未提交、日志不可用"，而不是含糊的"结果未知"。
    assert result["status"] == "journal_unavailable"
    assert result["journal_error"]
    assert result["next_action"] == "fix_journal"
    assert result["uncertain"] is False, "没有发出任何请求，订单结果并不未知"
    assert any("未决日志写入失败" in line for line in logs)


def test_run_sss_job_sends_no_post_when_the_lock_is_unavailable(
        tmp_path, monkeypatch):
    """拿不到批次独占锁（目录不可写 / 别的进程在跑）时同样零 POST。"""
    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", sss_max_workers=1)
    client = _FakeClient()
    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: client)

    def boom(*_a, **_k):
        from app.wps_atomicio import AtomicWriteError
        raise AtomicWriteError("锁文件不可写")

    monkeypatch.setattr(sss_module, "batch_submission_lock", boom)
    logs: list[str] = []

    result = run_sss_job(cfg, threading.Event(), progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")

    assert client.posts == []
    assert result["status"] == "journal_unavailable"
    assert any("无法取得下单批次互斥锁" in line for line in logs)


# ======================================================================
# A13b：inflight 残留（进程被杀）→ 下次运行阻断
# ======================================================================

def test_inflight_record_from_a_killed_process_blocks_the_next_run(tmp_path, monkeypatch):
    """POST 前写下的 inflight 如果一直没结案，就说明进程没能收尾 —— 必须继续阻断。"""
    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", sss_max_workers=1)
    journal_path = Path(cfg.sss_uncertain_path)
    day = sss_module.expected_delivery_date().isoformat()
    append_records(journal_path, batch_key(day, "", cfg.sss_account),
                   [_entry(status="inflight")], meta=_meta(delivery_date=day))

    state = blocking_state(journal_path, delivery_date=day,
                           account=cfg.sss_account, origin="https://sss.example.com")
    assert state["blocked"] is True and state["code"] == "unresolved_batch"

    client = _FakeClient()
    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: client)
    logs: list[str] = []
    result = run_sss_job(cfg, threading.Event(), progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")

    assert result["status"] == "blocked_by_uncertain"
    assert client.posts == []


# ======================================================================
# A15b：换账号 / 换平台的**处置**必须被拒绝（快照作用域绑定）
# ======================================================================

@pytest.fixture
def reviewed(tmp_path: Path):
    """一条未决记录 + 一次"站内确认没有"的核对快照。"""
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="a")], meta=_meta())
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(path, delivery_date="2026-09-12", account="18758187837",
                        origin="https://sss.example.com", fetch_json=fetch)
    return path, _id_of(path, "a")


def test_station_absent_refuses_after_the_account_changed(reviewed):
    path, record_id = reviewed
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="13900000000",
        decision="station_absent", record_ids=[record_id],
        confirm="station_absent", note="我已确认站内没有", origin="https://sss.example.com")

    assert got["ok"] is False
    assert got["code"] in ("cross_scope_record", "review_stale")
    assert "另一个账号" in got["reason"] or "当前账号" in got["reason"]


def test_station_absent_refuses_after_the_platform_changed(reviewed):
    """换网址之后，旧平台的核对结果对新平台无效（那边根本没查过订单）。"""
    path, record_id = reviewed
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[record_id],
        confirm="station_absent", note="我已确认站内没有",
        origin="https://other.example.com")

    assert got["ok"] is False
    assert got["code"] in ("cross_scope_record", "review_stale")


def test_station_absent_still_works_in_the_same_scope(reviewed):
    """反向锁：作用域一致时必须能正常解除（否则功能不可用）。"""
    path, record_id = reviewed
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[record_id],
        confirm="station_absent", note="我已确认站内没有",
        origin="https://sss.example.com")

    assert got["ok"] is True and got["changed"] is True
    state = blocking_state(path, delivery_date="2026-09-12", account="18758187837",
                           origin="https://sss.example.com")
    assert state["blocked"] is False


# ======================================================================
# A16b：scan_failed 不能支撑 station_absent，且仍保持阻断
# ======================================================================

def test_scan_failed_cannot_unblock(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="a")], meta=_meta())

    def boom(_path: str):
        raise RuntimeError("接口 500")

    snapshot = review.start_review(path, delivery_date="2026-09-12",
                                  account="18758187837",
                                  origin="https://sss.example.com", fetch_json=boom)
    assert snapshot["counts"]["scan_failed"] == 1

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[_id_of(path, "a")],
        confirm="station_absent", note="读取失败但我认为没有",
        origin="https://sss.example.com")

    assert got["ok"] is False and got["code"] == "station_not_absent"
    assert "查询失败或找到相似订单都不能当成没下单" in got["reason"]
    state = blocking_state(path, delivery_date="2026-09-12", account="18758187837",
                           origin="https://sss.example.com")
    assert state["blocked"] is True, "读取失败后必须继续阻断"


# ======================================================================
# A17b：核对快照过期 → 拒绝解除
# ======================================================================

def test_expired_review_snapshot_cannot_unblock(reviewed):
    path, record_id = reviewed
    snapshot = review.load_snapshot(path)
    assert snapshot is not None
    # 把快照时间推到 2 小时前（TTL 是 600 秒）
    stale = (dt.datetime.now() - dt.timedelta(hours=2)).isoformat(timespec="seconds")
    snapshot["created_at"] = stale
    review._atomic_write(review.review_snapshot_path(path), snapshot)

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[record_id],
        confirm="station_absent", note="我用的是旧证据",
        origin="https://sss.example.com")

    assert got["ok"] is False and got["code"] == "review_stale"
    assert "已过期" in got["reason"]


def test_record_not_in_review_is_rejected(reviewed):
    """记录在核对之后新增 → 不在快照里 → 不能拿别人的证据解除它。"""
    path, _record_id = reviewed
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="b", name="新人", phone="13900000002")],
                   meta=_meta())
    new_id = _id_of(path, "b")

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[new_id],
        confirm="station_absent", note="我认为也没有",
        origin="https://sss.example.com")

    assert got["ok"] is False
    assert got["code"] in ("review_stale", "record_not_in_review")


# ======================================================================
# A18 补强：keep 之后必须仍然阻断
# ======================================================================

def test_keep_keeps_blocking_and_never_writes(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="a")], meta=_meta())

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="keep", record_ids=[_id_of(path, "a")], confirm="keep",
        note="先不动", origin="https://sss.example.com")

    assert got["ok"] is True and got["status"] == "kept"
    assert got["changed"] is False
    assert got.get("cloud_write") is False
    state = blocking_state(path, delivery_date="2026-09-12", account="18758187837",
                           origin="https://sss.example.com")
    assert state["blocked"] is True


# ======================================================================
# A19：浏览器（Playwright）路径 —— POST 前留痕 + 异常归类与接口路径一致
# ======================================================================

class _FakePage:
    """只实现 sss.py 浏览器分支用到的那几个方法。"""

    def __init__(self, *, list_records=None, post_result=None,
                 post_raises: BaseException | None = None,
                 post_http: int = 200):
        self.calls: list[tuple[str, str]] = []
        self.posts: list[dict] = []
        self._list_records = list_records or []
        self._post_result = post_result
        self._post_raises = post_raises
        self._post_http = post_http

    # main.py 里用到的页面动作（本测试不走到）
    def goto(self, *_a, **_k) -> None: ...
    def locator(self, *_a, **_k):
        raise AssertionError("本测试不应调用页面元素定位")
    def wait_for_selector(self, *_a, **_k) -> None: ...
    def wait_for_timeout(self, *_a, **_k) -> None: ...
    def fill(self, *_a, **_k) -> None: ...

    def evaluate(self, _js: str, arg: dict) -> str:
        method = str(arg.get("method") or "")
        path = str(arg.get("path") or "")
        body = arg.get("body")
        self.calls.append((method, path))
        if method == "POST":
            self.posts.append(body or {})
            if self._post_raises is not None:
                raise self._post_raises
            payload = (self._post_result
                       if self._post_result is not None else {"success": True})
            if payload.get("success") and 200 <= int(self._post_http) < 300:
                # 站内真的落单了：把它加进列表，后续只读对账才可能确认。
                address = (body or {}).get("receiveAddress") or {}
                self._list_records = [*self._list_records, {
                    "recipientName": (body or {}).get("receiveName"),
                    "recipientPhone": [(body or {}).get("receivePhone")],
                    "recipientAddress": f"{address.get('addressDetail', '')}"
                                        f"{address.get('doorNum', '')}",
                    "expectedDeliveryTime": (body or {}).get("expectedDeliveryTime"),
                    "status": 2,
                }]
            return json.dumps({"http": self._post_http,
                               "text": json.dumps(payload, ensure_ascii=False)})
        if "list" in path:
            payload = {"success": True, "result": {"records": self._list_records,
                                                  "total": len(self._list_records)}}
        elif "store" in path:
            payload = {"success": True, "result": [{"name": "一口轻食", "id": 7}]}
        elif "account" in path:
            payload = {"success": True, "result": {"totalAmount": 1000.0}}
        else:
            payload = {"success": True, "result": []}
        return json.dumps({"http": 200, "text": json.dumps(payload, ensure_ascii=False)})


@pytest.fixture
def browser_env(tmp_path, monkeypatch):
    """把 Playwright 整体替换成假实现，只保留 sss.py 的浏览器分支逻辑。"""
    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", api_mode=False,
                  sss_max_workers=1)

    page = _FakePage()

    class _FakeBrowser:
        def new_page(self):
            return page

        def new_browser_cdp_session(self):
            raise AssertionError("不应真的做 CDP")

        def close(self) -> None: ...

    class _FakePlaywright:
        def __enter__(self):
            return self

        def __exit__(self, *_exc) -> None: ...

    fake_sync_api = types.ModuleType("playwright.sync_api")
    fake_sync_api.sync_playwright = lambda: _FakePlaywright()
    fake_playwright = types.ModuleType("playwright")
    fake_playwright.sync_api = fake_sync_api
    monkeypatch.setitem(sys.modules, "playwright", fake_playwright)
    monkeypatch.setitem(sys.modules, "playwright.sync_api", fake_sync_api)

    monkeypatch.setattr(sss_module, "_launch_browser",
                        lambda playwright, headless: _FakeBrowser())
    monkeypatch.setattr(sss_module, "_ensure_logged_in",
                        lambda page_, account, password, stop, cb: None)
    monkeypatch.setattr(sss_module, "_minimize_window", lambda browser, cb: None)
    return cfg, page


def test_browser_path_writes_inflight_before_post(browser_env, tmp_path):
    """A19：浏览器路径也必须在 POST 之前留下 inflight 记录。"""
    cfg, page = browser_env
    journal_path = Path(cfg.sss_uncertain_path)
    seen_statuses: list[str] = []
    real_append = sss_module.append_records

    def spy_append(path, key, entries, **kwargs):
        if path == journal_path or str(path) == str(journal_path):
            seen_statuses.extend(str(entry.get("status") or "")
                                 for entry in entries)
        return real_append(path, key, entries, **kwargs)

    sss_module.append_records, original = spy_append, sss_module.append_records
    try:
        result = run_sss_job(cfg, threading.Event(), progress_callback=lambda _m: None,
                             password="pw", captcha_callback=lambda _img: "1234")
    finally:
        sss_module.append_records = original

    assert result["created"] == 2, "假实现里两单都应成功"
    assert page.posts, "浏览器路径应该真的发出 POST"
    assert "inflight" in seen_statuses, "POST 之前必须先写 inflight"
    assert seen_statuses.index("inflight") < len(seen_statuses), seen_statuses

    # POST 之后收尾对账在站内找到了这两单 → 记录被标记为 resolved（这正是设计语义：
    # 只有"站内确实已有"才算确认），闸门随之解除。
    records = load_journal(journal_path)["records"]
    assert records and all(record["status"] == "resolved" for record in records), (
        [record["status"] for record in records])
    state = blocking_state(journal_path,
                           delivery_date=sss_module.expected_delivery_date().isoformat(),
                           account=cfg.sss_account, origin="https://sss.example.com")
    assert state["blocked"] is False


def test_browser_path_transport_error_is_unknown_not_failure(browser_env):
    """A19：浏览器提交遇到网络错误 → 必须留 unresolved，不能当成"明确拒绝"。"""
    cfg, page = browser_env
    page._post_raises = RuntimeError("Target page, context or browser has been closed")

    run_sss_job(cfg, threading.Event(), progress_callback=lambda _m: None,
                password="pw", captcha_callback=lambda _img: "1234")

    records = load_journal(Path(cfg.sss_uncertain_path))["records"]
    assert records, "网络错误后必须留下未决记录"
    assert all(record["status"] == "unresolved" for record in records), (
        "传输层错误绝不能被记成「服务端拒绝」——那样下一次运行就能重复下单")
    state = blocking_state(Path(cfg.sss_uncertain_path),
                           delivery_date=sss_module.expected_delivery_date().isoformat(),
                           account=cfg.sss_account, origin="https://sss.example.com")
    assert state["blocked"] is True


def test_browser_path_gateway_error_is_unknown(browser_env):
    """网关 5xx 同样是"没拿到可确认响应"，不能当拒绝。"""
    cfg, page = browser_env
    page._post_http = 502

    run_sss_job(cfg, threading.Event(), progress_callback=lambda _m: None,
                password="pw", captcha_callback=lambda _img: "1234")

    records = load_journal(Path(cfg.sss_uncertain_path))["records"]
    assert records and all(record["status"] == "unresolved" for record in records)


def test_browser_path_explicit_rejection_closes_the_record(browser_env):
    """反向锁：服务端明确 success:false 时可以安全关闭记录（否则永远卡住）。"""
    cfg, page = browser_env
    page._post_result = {"success": False, "message": "商品已下架"}

    run_sss_job(cfg, threading.Event(), progress_callback=lambda _m: None,
                password="pw", captcha_callback=lambda _img: "1234")

    records = load_journal(Path(cfg.sss_uncertain_path))["records"]
    assert records and all(record["status"] == "discarded" for record in records)


def test_missing_success_field_is_treated_as_unknown():
    """响应里没有 success 字段（协议变了/网关 JSON）→ 未知，不是拒绝。"""
    from app.sss import _check_success, _SubmissionUncertain

    with pytest.raises(_SubmissionUncertain):
        _check_success({"code": 0, "msg": "ok"})
    with pytest.raises(_SubmissionUncertain):
        _check_success({})


def test_explicit_false_is_a_definite_rejection():
    """显式 success=false 仍然是"服务端拒绝"（可安全关闭记录）。"""
    from app.sss import _check_success

    with pytest.raises(LookupError) as excinfo:
        _check_success({"success": False, "message": "商品已下架"})
    assert not isinstance(excinfo.value, sss_module._SubmissionUncertain)


# ======================================================================
# 操作槽位：完成态不能继续显示 active（§2.3）
# ======================================================================

def test_finished_operations_never_report_active(tmp_path, monkeypatch):
    """业务状态词可以是任何值；终结后 `active` 必须为假（前端据此解禁按钮）。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"x")
    bridge._config.excel_path = excel
    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})

    class _Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: _Cli())
    bridge.wps_preview()

    status = bridge.operation_status()
    assert status["active"] is False, "preview_ready 是完成态，不能继续显示进行中"
    assert status["operation"] is None
    assert status["last"]["status"] == "preview_ready"
    assert status["last"]["active"] is False


def test_unknown_business_status_is_still_terminal(tmp_path):
    """任意业务状态（这里用 sss 的 no_orders）终结后都不能停留在 active。"""
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    operation, conflict = bridge._reserve("sss", title="闪时送下单")
    assert conflict is None
    bridge._worker_operation = operation
    bridge._finish_task("闪时送没有需要下单的订单", {"status": "no_orders",
                                                    "processed": 0, "created": 0})

    status = bridge.operation_status()
    assert status["active"] is False
    assert status["last"]["status"] == "no_orders"
    assert status["last"]["active"] is False


def test_conflict_reports_the_running_operation(tmp_path):
    """冲突时必须能说明"谁在跑、下一步做什么"。"""
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    bridge._config.excel_path = tmp_path / "排单.xlsx"
    bridge._config.excel_path.write_bytes(b"x")
    operation, _ = bridge._reserve("order", title="订单处理", phase="登录")
    try:
        got = bridge.sss_uncertain_records()
        assert got["ok"] is True, "只读查询不该被互斥挡住"
        conflict = bridge.wps_preview()
        assert conflict["ok"] is False
        assert conflict["code"] == "operation_conflict"
        assert "订单处理" in conflict["reason"]
        assert conflict["next_action"]
        assert conflict["conflicting_operation"]["mode"] == "order"
    finally:
        bridge._operations.finish(operation, status="success")


# ======================================================================
# A23：跨范围与历史未决必须可见、可在正确范围核对
# ======================================================================

def test_a23_cross_scope_records_stay_visible(tmp_path, monkeypatch):
    """换账号/换日期之后，旧记录不能消失：要在 other_scope / history 里看得见。"""
    from app.bridge import Bridge

    # 日期必须用程序真正会用的"送达日"：其它日期会被归到 history。
    from app.sss import expected_delivery_date

    day = expected_delivery_date().isoformat()
    yesterday = (expected_delivery_date() - dt.timedelta(days=1)).isoformat()
    state = tmp_path / "sss_uncertain.json"
    append_records(state, batch_key(day, "", "18758187837"),
                   [_entry(identifier="今天的单")], meta=_meta(delivery_date=day))
    append_records(state, batch_key(yesterday, "", "18758187837"),
                   [_entry(identifier="昨天的单", phone="13900000009")],
                   meta=_meta(delivery_date=yesterday))
    append_records(state, batch_key(day, "", "13900000000"),
                   [_entry(identifier="别的账号的单", phone="13900000008")],
                   meta=_meta(account="13900000000", delivery_date=day))

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.sss_account = "18758187837"
    bridge._config.sss_url = "https://sss.example.com/takeout"

    got = bridge.sss_uncertain_records()

    assert got["ok"] is True
    identifiers = {item["identifier"] for item in got["records"]}
    assert identifiers == {"今天的单"}, "当前范围只列自己的那一条"
    groups = got["groups"]
    assert {item["identifier"] for item in groups["other_scope"]} == {"别的账号的单"}
    assert {item["identifier"] for item in groups["history"]} == {"昨天的单"}
    assert got["other_scope_count"] == 1
    assert "其它范围" in (got["next_action"] or ""), got["next_action"]
    # 脱敏：列表里不能出现完整手机号
    blob = json.dumps(got, ensure_ascii=False, default=str)
    assert "13900000008" not in blob and "13900000009" not in blob
    # 只读：查询不该改动日志
    assert load_journal(state)["records"][0]["status"] == "unresolved"


def test_a23_cannot_resolve_a_record_from_another_scope(tmp_path):
    """用当前账号的证据处置别的账号的记录必须被拒绝。"""
    from app import sss_review as review

    state = tmp_path / "sss_uncertain.json"
    append_records(state, batch_key("2026-09-12", "", "13900000000"),
                   [_entry(identifier="别的账号的单", phone="13900000008")],
                   meta=_meta(account="13900000000"))
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(state, delivery_date="2026-09-12", account="13900000000",
                        origin="https://sss.example.com", fetch_json=fetch)
    record_id = _id_of(state, "别的账号的单")

    got = review.resolve_records(
        state, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[record_id],
        confirm="station_absent", note="我用当前账号确认过",
        origin="https://sss.example.com")

    assert got["ok"] is False and got["code"] == "cross_scope_record"
    state_now = blocking_state(state, delivery_date="2026-09-12",
                               account="13900000000",
                               origin="https://sss.example.com")
    assert state_now["blocked"] is True, "别的账号的记录不能被当前账号解除"


# ======================================================================
# A28：Bridge 全异常路径都不留僵死占位
# ======================================================================

def test_a28_bridge_releases_the_slot_on_internal_failures(tmp_path, monkeypatch):
    """保存配置/写密钥链/起线程失败后，占位必须释放，下一个操作仍可启动。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    excel = tmp_path / "排单.xlsx"
    _write_workbook(excel)

    payload = {
        "url": "https://m.icall.me/admin/#/login", "phone": "13800000000",
        "password": "pw", "excel": str(excel), "date": "", "count": "",
        "remember": True, "api_mode": True,
    }

    def boom_save(*_a, **_k):
        raise OSError("磁盘满")

    monkeypatch.setattr(type(bridge._config), "save", boom_save)
    got = bridge.start_order(payload)
    assert got["ok"] is False and got["reason"] == "internal_error"
    assert bridge.operation_status()["active"] is False, "失败后不能留下僵死占位"

    monkeypatch.undo()
    started: list[str] = []
    monkeypatch.setattr(bridge, "_launch",
                        lambda mode, cfg, count, password: started.append(mode) or True)
    assert bridge.start_order(payload)["ok"] is True
    assert started == ["order"], "上一次失败不能影响下一次启动"

    # 密钥链失败同样要释放
    bridge.operation_status()  # 触发一次读取（不改变状态）
    bridge._operations.finish(bridge._worker_operation, status="success")
    bridge._worker_operation = None

    def boom_password(*_a, **_k):
        raise RuntimeError("keyring 不可用")

    monkeypatch.setattr(bridge_module, "set_password", boom_password)
    got = bridge.start_order(payload)
    assert got["ok"] is False and got["reason"] == "internal_error"
    assert bridge.operation_status()["active"] is False


def test_a28_conflict_does_not_disturb_the_running_operation(tmp_path):
    """冲突返回不能影响正在跑的那个操作（它的槽位与状态必须原样保留）。"""
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    bridge._config.excel_path = _write_workbook(tmp_path / "排单.xlsx")

    operation, _conflict = bridge._reserve("order", title="订单处理", phase="登录")
    try:
        got = bridge.check_updates()
        assert got["ok"] is False and got["reason"] == "operation_conflict"
        status = bridge.operation_status()
        assert status["active"] is True
        assert status["operation"]["mode"] == "order"
        assert status["operation"]["phase"] == "登录"
    finally:
        bridge._operations.finish(operation, status="success")


# ======================================================================
# A29：只有样式变化 / 同尺寸同 mtime 的改动也作废令牌
# ======================================================================

def test_a29_style_only_change_invalidates_the_token(tmp_path, monkeypatch):
    """只改单元格格式（数据完全一样）也要让预览令牌失效。"""
    from openpyxl import load_workbook

    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = _write_workbook(tmp_path / "排单.xlsx")
    bridge._config.excel_path = excel

    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})

    class Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: Cli())
    calls: list[str] = []
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: calls.append("apply") or {
                            "sheets": [], "written": 0, "failed": 0})

    preview = bridge.wps_preview()
    assert preview["ok"] is True

    # 只改字体（数据一个字节都不变），并保持文件大小/mtime 可见性无关紧要
    workbook = load_workbook(excel)
    sheet = workbook["东湖中餐"]
    sheet["B3"].font = sheet["B3"].font.copy(bold=True)
    workbook.save(excel)

    got = bridge.wps_upload(preview["preview_id"])
    assert got["ok"] is False and got["code"] == "preview_changed"
    assert "local_file" in (got.get("changed") or [])
    assert calls == [], "样式变了也不能上传"


def test_a29_same_size_change_is_caught_by_the_content_hash(tmp_path, monkeypatch):
    """同尺寸改动（内容变了但长度相同）同样要被内容哈希挡住。"""
    from app import bridge as bridge_module
    from app.bridge import Bridge

    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge._config.wps_enabled = True
    excel = tmp_path / "排单.xlsx"
    excel.write_bytes(b"A" * 4096)
    bridge._config.excel_path = excel

    monkeypatch.setattr(bridge_module, "effective_tables",
                        lambda _cfg: {"东湖中餐": {"file_id": "F1"}})
    monkeypatch.setattr(bridge_module, "read_local_orders_from_bytes",
                        lambda data, log=None: {"东湖中餐": []})
    monkeypatch.setattr(bridge_module, "build_plan", lambda cli, **kw: [])
    monkeypatch.setattr(bridge_module, "format_plan", lambda plans: "")
    monkeypatch.setattr(bridge_module, "summarize_plan", lambda plans: {})

    class Cli:
        path = "/fake/kdocs-cli"

        def authenticated(self) -> bool:
            return True

    monkeypatch.setattr(bridge, "_wps_cli", lambda: Cli())
    monkeypatch.setattr(bridge_module, "apply_plan",
                        lambda *a, **k: {"sheets": [], "written": 0, "failed": 0})

    preview = bridge.wps_preview()
    assert preview["ok"] is True

    same_size = b"B" * 4096          # 长度相同、内容不同
    excel.write_bytes(same_size)
    got = bridge.wps_upload(preview["preview_id"])
    assert got["ok"] is False and got["code"] == "preview_changed"


def _write_workbook(path):
    """写一个能被 openpyxl 解析的最小排单表。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "东湖中餐"
    ws.cell(2, 2, "姓名")
    ws.cell(3, 2, "老人")
    ws.cell(3, 3, "小")
    ws.cell(3, 4, "13800000000")
    ws.cell(3, 14, 6)
    wb.save(path)
    return path
