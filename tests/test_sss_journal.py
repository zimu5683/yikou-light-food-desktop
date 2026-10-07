"""闪时送未决日志：跨运行阻断、只读核对、人工处置与两条提交路径的接入。

这些是**最重要的安全属性**：闪时送接口没有客户端幂等键，重复 POST 会真的多
下一单（客户多收一份饭）。因此：

* POST 之前先落 ``inflight``，拿不到确定响应保持 ``unresolved``；
* 同批次再有未决记录时，任何运行都不许自动发 POST；
* 换名单来源（wps→excel）、重启程序、再点一次「开始下单」都不能绕过；
* 日志损坏时失败关闭（绝不当作"没有未决记录"）；
* 只有"站内确实已有这一单"或"人工核对证明站内没有"才能解除阻断。
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from app import sss_journal as journal
from app import sss_review as review
from app.sss import OrderFingerprint, _JournalHooks, run_sss_job
from app.sss_journal import (UncertainJournalError, append_records, batch_key,
                             blocking_state, discard_records,
                             journal_fingerprint, load_journal,
                             mask_contact, normalise_account,
                             pending_record_views, platform_origin,
                             resolve_records)


def _fingerprint(name="张", phone="13800000000", **overrides):
    base = {
        "receive_name": name, "receive_phone": phone, "door_num": "A1",
        "expected_delivery_time": "2026-09-12 11:00:00",
        "account": "18758187837", "store_id": "1", "goods_name": "轻食",
        "goods_num": "1", "address_detail": "浙江农林大学东湖校区",
        "area_code": "330110", "lnt": "119.7", "lat": "30.2", "order_type": "2",
    }
    base.update(overrides)
    return OrderFingerprint(**base)


def _entry(name="张", phone="13800000000", **overrides):
    # 先把"记录级字段"取出来，剩下的才是要盖进指纹的字段 —— 否则 status/error
    # 会被当成指纹参数传进 OrderFingerprint。
    identifier = overrides.pop("identifier", f"第 3 行 {name}")
    sheet = overrides.pop("sheet", "午餐")
    batch_id = overrides.pop("batch_id", "batch1")
    error = overrides.pop("error", "请求超时")
    status = overrides.pop("status", "unresolved")
    return {
        "identifier": identifier,
        "sheet": sheet,
        "batch_id": batch_id,
        "fingerprint": _fingerprint(name, phone, **overrides).as_dict(),
        "error": error,
        "status": status,
    }


def _meta(account="18758187837", platform="https://sss.example.com",
          delivery_date="2026-09-12", source="wps"):
    return {"account": account, "platform": platform,
            "delivery_date": delivery_date, "source": source,
            "batch_started_at": 1_700_000_000.0}


# ----------------------------------------------------------------------
# 规范化与身份
# ----------------------------------------------------------------------

def test_normalise_account_only_merges_whitespace_and_width():
    assert normalise_account(" 187 5818 7837 ") == "18758187837"
    assert normalise_account("１８７５８１８７８３７") == "18758187837"   # 全角
    assert normalise_account("01875818783") == "01875818783", "不能丢前导零"
    assert normalise_account(None) == ""


def test_batch_key_ignores_the_order_source():
    """换名单来源（wps/excel）不能解除同一批次的未决记录。"""
    assert batch_key("2026-09-12", "wps", "187 5818 7837") == \
        batch_key("2026-09-12", "excel", "18758187837")
    assert batch_key("2026-09-12", "", "a") != batch_key("2026-09-13", "", "a")


def test_platform_origin_normalises_equivalent_urls():
    same = platform_origin(url="HTTPS://SSS.Example.COM:443/takeout")
    assert same == platform_origin(url="https://sss.example.com/other/path")
    assert platform_origin(url="http://sss.example.com:8080/x") == \
        "http://sss.example.com:8080"


@pytest.mark.parametrize("bad", [
    "https://sss.example.com./takeout",     # 尾点：会被当成不同来源
    "sss.example.com",                       # 缺协议
    "https://中文.example.com/",             # 非 ASCII 域名
    "https://user:pw@sss.example.com/",      # 带凭据
    "https://sss.example.com:0/",            # 端口非法
    "https://sss.example.com:99999/",
    "https://sss.example.com/ takeout",
])
def test_platform_origin_rejects_non_canonical_urls(bad):
    """这些写法必须 fail-closed，绝不能为它们另建一个作用域。"""
    with pytest.raises(UncertainJournalError):
        platform_origin(url=bad)


# ----------------------------------------------------------------------
# 写入与阻断
# ----------------------------------------------------------------------

def test_unresolved_record_blocks_the_same_batch(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    append_records(path, batch_key("2026-09-12", "", "18758187837"),
                   [_entry()], meta=_meta())

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="187 5818 7837",
                           origin="https://sss.example.com")

    assert state["blocked"] is True and state["code"] == "unresolved_batch"
    assert len(state["same_scope"]) == 1
    assert "重复提交会真的多下单" in state["reason"]


def test_switching_source_does_not_unblock(tmp_path: Path):
    """用旧键（三段）写入的记录，换 source 后仍然阻断。"""
    path = tmp_path / "sss_uncertain.json"
    append_records(path, journal.legacy_batch_key("2026-09-12", "wps", "18758187837"),
                   [_entry()], meta=_meta(source="wps"))

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837",
                           origin="https://sss.example.com")
    assert state["blocked"] is True


def test_switching_account_blocks_as_cross_scope(tmp_path: Path):
    """换账号：旧记录不算当前批次，但也不能静默丢弃。"""
    path = tmp_path / "sss_uncertain.json"
    append_records(path, batch_key("2026-09-12", "", "18758187837"),
                   [_entry()], meta=_meta(account="18758187837"))

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="13900000000",
                           origin="https://sss.example.com")

    assert state["blocked"] is True and state["code"] == "cross_scope_unresolved"
    assert len(state["cross_scope"]) == 1
    assert "账号或平台地址与当前不同" in state["reason"]


def test_switching_platform_origin_blocks_as_cross_scope(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    append_records(path, batch_key("2026-09-12", "", "18758187837"),
                   [_entry()], meta=_meta())

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837",
                           origin="https://other.example.com")
    assert state["blocked"] is True and state["code"] == "cross_scope_unresolved"


def test_another_day_does_not_block_today_but_stays_visible(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    append_records(path, batch_key("2026-09-11", "", "18758187837"),
                   [_entry()], meta=_meta(delivery_date="2026-09-11"))

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837",
                           origin="https://sss.example.com")
    assert state["blocked"] is False
    assert len(state["other_pending"]) == 1, "旧记录不能被静默丢弃"


def test_corrupt_journal_fails_closed(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    path.write_text("{ 不是 json", encoding="utf-8")
    with pytest.raises(UncertainJournalError):
        load_journal(path)
    with pytest.raises(UncertainJournalError):
        blocking_state(path, delivery_date="2026-09-12", account="a",
                       origin="https://sss.example.com")


def test_journal_with_unknown_status_is_rejected(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    path.write_text(json.dumps({"version": 1, "records": [
        {"journal_id": "x", "status": "unknown-status"}]}), encoding="utf-8")
    with pytest.raises(UncertainJournalError):
        load_journal(path)


def test_resolve_and_discard_close_records(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="a"), _entry(identifier="b", phone="13900000001")],
                   meta=_meta())
    before = journal_fingerprint(path)

    assert resolve_records(path, key, ["a"], note="站内已有") == 1
    assert journal_fingerprint(path) != before, "状态变化必须改变指纹"
    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837", origin="https://sss.example.com")
    assert state["blocked"] is True and len(state["same_scope"]) == 1

    assert discard_records(path, key, ["b"], reason="明确未派发") == 1
    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837", origin="https://sss.example.com")
    assert state["blocked"] is False


def test_pending_views_mask_contact_details(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    append_records(path, batch_key("2026-09-12", "", "18758187837"),
                   [_entry()], meta=_meta())
    views = pending_record_views(path, batch_key("2026-09-12", "", "18758187837"))
    assert views["counts"]["active"] == 1
    record = views["records"][0]
    assert record["phone"] == "138****0000"
    assert record["account"] == "187****7837"
    assert "13800000000" not in json.dumps(views, ensure_ascii=False)


def test_mask_contact_keeps_shape_for_other_values():
    assert mask_contact("") == ""
    assert mask_contact("abc") == "***"
    assert mask_contact("18758187837") == "187****7837"


def test_append_records_keeps_other_active_records(tmp_path: Path):
    """另一轮的未决记录不能被本次追加覆盖掉（否则跨运行阻断会漏单）。"""
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="旧的", name="旧")], meta=_meta())
    append_records(path, key, [_entry(identifier="新的", name="新", phone="13900000002")],
                   meta=_meta())
    records = load_journal(path)["records"]
    identifiers = {record["identifier"] for record in records}
    assert identifiers == {"旧的", "新的"}


# ----------------------------------------------------------------------
# _JournalHooks：提交前落意图、提交后定状态
# ----------------------------------------------------------------------

def _hooks(path, tmp_path):
    return _JournalHooks(path, key=batch_key("2026-09-12", "", "18758187837"),
                         meta=_meta())


def _task(identifier: str) -> dict:
    """与 run_sss_job 里真实任务同构的最小任务。"""
    return {"identifier": identifier, "sheet": "午餐", "batch_id": "b",
            "fingerprint": _fingerprint(), "account": "18758187837",
            "client_request_id": ""}


def test_prepare_fails_closed_when_journal_is_unwritable(tmp_path: Path):
    """留痕失败就必须停止：发出去但没留痕是重复下单的根源。"""
    hooks = _hooks(tmp_path / "nope" / "dir" / "sss_uncertain.json", tmp_path)
    # 目录无法创建（父路径是文件）
    blocker = tmp_path / "nope"
    blocker.write_text("x", encoding="utf-8")
    assert hooks.prepare([_task("第 3 行 张")]) is False
    assert "未决日志写入失败" in hooks.error


def test_prepare_then_finalize_marks_unsent_as_discarded(tmp_path: Path):
    from app.sss import _SubmitResult

    path = tmp_path / "sss_uncertain.json"
    hooks = _hooks(path, tmp_path)
    tasks = [_task("a"), _task("b"), _task("c")]
    assert hooks.prepare(tasks) is True
    record_ids = {record["identifier"] for record in load_journal(path)["records"]}
    assert record_ids == {"a", "b", "c"}
    assert all(record["status"] == "inflight"
               for record in load_journal(path)["records"])

    result = _SubmitResult()
    result.succeeded.add("a")
    result.failures.append(("b", "余额不足"))
    result.uncertain.append(("c", "请求超时"))
    assert hooks.finalize(tasks, result) is True

    by_id = {record["identifier"]: record for record in load_journal(path)["records"]}
    assert by_id["a"]["status"] == "unresolved", (
        "返回 success 也要等站内对账确认，不能直接删记录")
    assert by_id["b"]["status"] == "discarded"
    assert by_id["c"]["status"] == "unresolved"


# ----------------------------------------------------------------------
# 只读核对与人工处置
# ----------------------------------------------------------------------

def _fetch_queue(*payloads):
    items = list(payloads)

    def fetch_json(path: str):
        if not items:
            return {"success": True, "result": {"records": [], "total": 0}}
        return items.pop(0)
    return fetch_json


def _order_record(name="张", phone="13800000000"):
    return {
        "recipientName": name, "recipientPhone": [phone],
        "recipientAddress": "浙江农林大学东湖校区A1",
        "expectedDeliveryTime": "2026-09-12 11:00:00",
        "status": 2,
    }


@pytest.fixture
def journaled(tmp_path: Path):
    path = tmp_path / "sss_uncertain.json"
    key = batch_key("2026-09-12", "", "18758187837")
    append_records(path, key, [_entry(identifier="a"), _entry(identifier="b", phone="13900000001")],
                   meta=_meta())
    return path, key


def _id_of(path, identifier: str) -> str:
    """日志记录的 journal_id（= 提交前由批次+行位+指纹派生）。"""
    for record in load_journal(path)["records"]:
        if record["identifier"] == identifier:
            return record["journal_id"]
    raise AssertionError(f"日志里没有 {identifier}")


def test_review_confirms_matched_records_and_keeps_the_rest(journaled):
    path, key = journaled
    fetch = _fetch_queue({"success": True, "result": {
        "records": [_order_record()], "total": 1}})

    snapshot = review.start_review(path, delivery_date="2026-09-12",
                                  account="18758187837", fetch_json=fetch)

    assert snapshot["counts"]["station_confirmed"] == 1
    classifications = {item["journal_id"]: item["classification"]
                       for item in snapshot["results"]}
    assert classifications[_id_of(path, "a")] == "station_confirmed"
    assert classifications[_id_of(path, "b")] == "station_missing"

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837", origin="https://sss.example.com")
    assert state["blocked"] is True, "仍有一条站内没有 → 保持阻断"
    assert len(state["same_scope"]) == 1


def test_review_classifies_scan_failure_separately(journaled):
    """查询失败绝不能当成"站内没有订单"。"""
    path, _key = journaled

    def boom(_path: str):
        raise RuntimeError("接口 500")

    snapshot = review.start_review(path, delivery_date="2026-09-12",
                                  account="18758187837", fetch_json=boom)

    assert snapshot["counts"]["scan_failed"] == 2
    assert all(item["classification"] == review.SCAN_FAILED
               for item in snapshot["results"])
    assert "无法判定" in snapshot["results"][0]["reason"]


def test_review_does_not_cross_platform_scope(journaled):
    path, key = journaled
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})

    snapshot = review.start_review(
        path, delivery_date="2026-09-12", account="18758187837",
        origin="https://another.example.com", fetch_json=fetch)

    assert snapshot["results"] == []
    assert snapshot["platform_origin"] == "https://another.example.com"
    blocked = blocking_state(
        path, delivery_date="2026-09-12", account="18758187837",
        origin="https://another.example.com")
    assert blocked["blocked"] is True
    assert blocked["cross_scope"]

    resolved = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        origin="https://another.example.com", decision="station_present",
        record_ids=[_id_of(path, "a")], confirm="station_present",
        note="当前平台不应处置旧平台记录")
    assert resolved["ok"] is False
    assert resolved["code"] == "cross_scope_record"


def test_station_absent_requires_fresh_evidence_and_all_missing(journaled):
    path, key = journaled
    ids = [_id_of(path, "a"), _id_of(path, "b")]
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(path, delivery_date="2026-09-12",
                        account="18758187837", fetch_json=fetch)

    no_confirm = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=ids, confirm="station_present",
        note="已打电话确认站内没有")
    assert no_confirm["ok"] is False and no_confirm["code"] == "confirm_mismatch"

    short_note = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=ids, confirm="station_absent",
        note="嗯")
    assert short_note["ok"] is False and short_note["code"] == "note_too_short"

    partial = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[ids[0]], confirm="station_absent",
        note="只确认了第一条")
    assert partial["ok"] is True and partial["changed"] is True
    assert partial["cloud_write"] is False
    assert "下一次点「开始下单」才会真正下单" in partial["reason"]

    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837", origin="https://sss.example.com")
    assert state["blocked"] is True, "未处置的那条仍要阻断"

    # 上一次处置已经改动了日志 → 旧证据作废，必须**重新核对**才能再处置。
    stale = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[ids[1]], confirm="station_absent",
        note="第二条也确认没有")
    assert stale["ok"] is False and stale["code"] == "review_stale"

    fetch2 = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(path, delivery_date="2026-09-12",
                        account="18758187837", fetch_json=fetch2)
    rest = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[ids[1]], confirm="station_absent",
        note="第二条也确认没有")
    assert rest["ok"] is True
    state = blocking_state(path, delivery_date="2026-09-12",
                           account="18758187837", origin="https://sss.example.com")
    assert state["blocked"] is False


def test_station_absent_refuses_without_snapshot(journaled):
    path, _key = journaled
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=[_id_of(path, "a")],
        confirm="station_absent", note="我确定没有")
    assert got["ok"] is False and got["code"] == "review_required"


def test_station_absent_refuses_when_snapshot_lists_another_class(journaled):
    """证据里只要有"站内已找到"的记录，就不能整体按"站内没有"解除。"""
    path, key = journaled
    ids = [_id_of(path, "a"), _id_of(path, "b")]
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(path, delivery_date="2026-09-12",
                        account="18758187837", fetch_json=fetch)
    # 手工把快照的第一条改成 station_confirmed（模拟"自动确认没落盘"的边界情形：
    # 记录仍是激活的，但证据说站内已经有它）。
    snapshot = review.load_snapshot(path)
    assert snapshot is not None
    snapshot["results"][0]["classification"] = review.STATION_CONFIRMED
    snapshot["journal_fingerprint"] = journal_fingerprint(path)
    review._atomic_write(review.review_snapshot_path(path), snapshot)

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=ids, confirm="station_absent",
        note="我确定没有")
    assert got["ok"] is False and got["code"] == "station_not_absent"
    assert "查询失败或找到相似订单都不能当成没下单" in got["reason"]


def test_station_absent_refuses_when_journal_changed_after_review(journaled):
    path, key = journaled
    fetch = _fetch_queue({"success": True, "result": {"records": [], "total": 0}})
    review.start_review(path, delivery_date="2026-09-12",
                        account="18758187837", fetch_json=fetch)
    ids = [_id_of(path, "a"), _id_of(path, "b")]
    # 核对之后日志又被改过（例如另一个窗口又跑了一次）
    append_records(path, key, [_entry(identifier="c", name="新人",
                                      phone="13900000003")], meta=_meta())

    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_absent", record_ids=ids, confirm="station_absent",
        note="我确定没有")
    assert got["ok"] is False and got["code"] == "review_stale"


def test_station_present_resolves_records(journaled):
    path, _key = journaled
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="station_present", record_ids=[_id_of(path, "a")],
        confirm="station_present", note="站内确实有这一单")
    assert got["ok"] is True and got["changed"] is True
    assert got["cloud_write"] is False


def test_keep_is_a_noop(journaled):
    path, _key = journaled
    before = load_journal(path)
    ids = [_id_of(path, "a"), _id_of(path, "b")]
    got = review.resolve_records(
        path, delivery_date="2026-09-12", account="18758187837",
        decision="keep", record_ids=ids, confirm="keep", note="先不动")
    assert got["ok"] is True and got["status"] == "kept"
    assert got["changed"] is False
    assert load_journal(path) == before, "keep 不能改动任何记录"


# ----------------------------------------------------------------------
# run_sss_job：阻断与两条提交路径都留痕
# ----------------------------------------------------------------------

class _FakeClient:
    """SssApiClient 替身：登录、门店/地址、余额、列表、下单。"""

    def __init__(self, *, order_fail: bool = False, order_timeout: bool = False,
                 list_payload=None, login_payload=None):
        self.order_fail = order_fail
        self.order_timeout = order_timeout
        self.posts: list[dict] = []
        self._list_payload = list_payload or {"success": True, "result": {
            "records": [], "total": 0}}
        self._login_payload = login_payload or {"success": True}

    def fetch_captcha(self) -> bytes:
        return b"png"

    def login(self, _code: str) -> dict:
        return self._login_payload

    def get_json(self, path: str) -> dict:
        if "list" in path:
            payload = self._list_payload
            self._list_payload = {"success": True, "result": {"records": [], "total": 0}}
            return payload
        if "store" in path:
            return {"success": True, "result": [{"name": "一口轻食", "id": 7}]}
        if "account" in path:
            return {"success": True, "result": {"totalAmount": 1000.0}}
        return {"success": True, "result": []}

    def post_json(self, _path: str, payload: dict) -> dict:
        self.posts.append(payload)
        if self.order_timeout:
            from app.api_client import SssTransportError
            raise SssTransportError("读取超时")
        if self.order_fail:
            return {"success": False, "message": "商品已下架"}
        return {"success": True}

    def fork(self):
        return self

    def close(self) -> None:
        pass


def _config(tmp_path, **overrides):
    from app.config import AppConfig
    cfg = AppConfig(config_path=str(tmp_path / "config.json"))
    cfg.sss_dry_run = False
    cfg.sss_preflight = False
    cfg.api_mode = True
    cfg.sss_account = "18758187837"
    # 固定地址：测试只关心提交/留痕，不关心地址解析
    cfg.sss_use_fixed_address = True
    cfg.sss_url = "https://sss.example.com/takeout"
    cfg.sss_excel_path = tmp_path / "闪时送.xlsx"
    cfg.sss_uncertain_path = str(tmp_path / "sss_uncertain.json")
    for key, value in overrides.items():
        setattr(cfg, key, value)
    return cfg


def _write_sss_excel(tmp_path, rows=2):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.cell(2, 1, "姓名")
    for index in range(rows):
        ws.cell(3 + index, 1, f"客户{index}")
        ws.cell(3 + index, 2, f"A{index + 1}")
        ws.cell(3 + index, 3, f"138000000{index:02d}")
    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    return path


def test_run_sss_job_is_blocked_by_unresolved_records(tmp_path, monkeypatch):
    """未决记录存在时：不登录、不发任何下单请求。"""
    import app.sss as sss_module

    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel")
    journal_path = Path(cfg.sss_uncertain_path)
    append_records(journal_path, batch_key(
        __import__("app.sss", fromlist=["x"]).expected_delivery_date().isoformat(),
        "", cfg.sss_account), [_entry()], meta=_meta(
            delivery_date=__import__("app.sss", fromlist=["x"]).expected_delivery_date().isoformat()))

    created: list = []

    def fake_client(*args, **kwargs):
        created.append(True)
        return _FakeClient()

    monkeypatch.setattr(sss_module, "SssApiClient", fake_client)
    logs: list[str] = []
    result = run_sss_job(cfg, __import__("threading").Event(),
                         progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")

    assert result["status"] == "blocked_by_uncertain"
    assert result["blocked"] is True and result["created"] == 0
    assert created == [], "阻断时连客户端都不该创建（不会登录）"
    assert any("重复提交会真的多下单" in line for line in logs)


def test_run_sss_job_refuses_when_journal_is_corrupt(tmp_path, monkeypatch):
    import app.sss as sss_module

    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel")
    Path(cfg.sss_uncertain_path).write_text("{ 坏了", encoding="utf-8")

    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: _FakeClient())
    logs: list[str] = []
    result = run_sss_job(cfg, __import__("threading").Event(),
                         progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")
    assert result["status"] == "uncertain_journal_unreadable"
    assert result["created"] == 0


def test_run_sss_job_records_unresolved_on_timeout(tmp_path, monkeypatch):
    """POST 超时 → 留下 unresolved 记录，且下一次运行被阻断。"""
    import app.sss as sss_module

    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", sss_max_workers=1)
    client = _FakeClient(order_timeout=True)
    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: client)

    logs: list[str] = []
    result = run_sss_job(cfg, __import__("threading").Event(),
                         progress_callback=logs.append,
                         password="pw", captcha_callback=lambda _img: "1234")

    records = load_journal(Path(cfg.sss_uncertain_path))["records"]
    assert records, "超时后必须留下未决记录"
    assert all(record["status"] == "unresolved" for record in records)
    assert result["uncertain"] is True
    assert result["uncertain_pending"] == 2, "两条超时记录都要留着"
    posts_after_first = len(client.posts)
    assert posts_after_first == 2

    # 再跑一次：直接被阻断，一个 POST 都不发
    second_logs: list[str] = []
    again = run_sss_job(cfg, __import__("threading").Event(),
                        progress_callback=second_logs.append,
                        password="pw", captcha_callback=lambda _img: "1234")
    assert again["status"] == "blocked_by_uncertain"
    assert len(client.posts) == posts_after_first, "第二次运行不能发出任何 POST"


def test_run_sss_job_discards_records_for_clear_rejection(tmp_path, monkeypatch):
    """服务端明确拒绝（success=false）→ 关闭记录，不会阻断下一次运行。"""
    import app.sss as sss_module

    _write_sss_excel(tmp_path)
    cfg = _config(tmp_path, sss_order_source="excel", sss_max_workers=1)
    client = _FakeClient(order_fail=True)
    monkeypatch.setattr(sss_module, "SssApiClient", lambda *a, **k: client)

    run_sss_job(cfg, __import__("threading").Event(), progress_callback=lambda _m: None,
                password="pw", captcha_callback=lambda _img: "1234")

    records = load_journal(Path(cfg.sss_uncertain_path))["records"]
    assert records and all(record["status"] == "discarded" for record in records)
    state = blocking_state(Path(cfg.sss_uncertain_path), delivery_date=(
        __import__("app.sss", fromlist=["x"]).expected_delivery_date().isoformat()),
        account=cfg.sss_account, origin="https://sss.example.com")
    assert state["blocked"] is False


def test_concurrent_batch_is_refused(tmp_path):
    """另一个执行体（线程/进程）不能同时跑同一批次：那会产生两次 POST。"""
    from concurrent.futures import ThreadPoolExecutor

    from app.sss_journal import batch_submission_lock

    path = tmp_path / "sss_uncertain.json"
    lock = batch_submission_lock(path, batch_key("2026-09-12", "", "a"))
    try:
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(batch_submission_lock,
                                 path, batch_key("2026-09-12", "", "a"))
            with pytest.raises(UncertainJournalError) as excinfo:
                future.result(timeout=10)
        assert "拒绝并发提交" in str(excinfo.value)
    finally:
        lock.release()
    # 释放后可以再次获取（跨线程）
    pool = ThreadPoolExecutor(max_workers=1)
    try:
        second = pool.submit(batch_submission_lock, path,
                             batch_key("2026-09-12", "", "a")).result(timeout=10)
    finally:
        pool.shutdown()
    second.release()
