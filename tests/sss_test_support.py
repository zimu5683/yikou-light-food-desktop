"""闪时送（sss）测试的公共替身与夹具。

原先这些 helper 住在 ``tests/test_sss_journal.py``（未决日志用例的宿主）。
2026-10-12 移除未决阻断机制后，日志用例整体删除，但仍被其它用例使用的
helper 迁到这里，避免"删一个文件顺带删掉别人的夹具"。
"""
from __future__ import annotations


class _FakeClient:
    """SssApiClient 替身：登录、门店/地址、余额、列表、下单。"""

    def __init__(self, *, order_fail: bool = False, order_timeout: bool = False,
                 list_payload=None, login_payload=None):
        self.order_fail = order_fail
        self.order_timeout = order_timeout
        self.posts: list[dict] = []
        self._list_payload = list_payload or {"success": True, "result": {
            "records": [], "total": 0}}
        self._login_payload = login_payload or {"success": True}

    def fetch_captcha(self) -> bytes:
        return b"png"

    def login(self, _code: str) -> dict:
        return self._login_payload

    def get_json(self, path: str) -> dict:
        if "list" in path:
            payload = self._list_payload
            self._list_payload = {"success": True, "result": {"records": [], "total": 0}}
            return payload
        if "store" in path:
            return {"success": True, "result": [{"name": "一口轻食", "id": 7}]}
        if "account" in path:
            return {"success": True, "result": {"totalAmount": 1000.0}}
        return {"success": True, "result": []}

    def post_json(self, _path: str, payload: dict) -> dict:
        self.posts.append(payload)
        if self.order_timeout:
            from app.api_client import SssTransportError
            raise SssTransportError("读取超时")
        if self.order_fail:
            return {"success": False, "message": "商品已下架"}
        return {"success": True}

    def fork(self):
        return self

    def close(self) -> None:
        pass


def _config(tmp_path, **overrides):
    from app.config import AppConfig
    cfg = AppConfig(config_path=str(tmp_path / "config.json"))
    cfg.sss_dry_run = False
    cfg.sss_preflight = False
    cfg.api_mode = True
    cfg.sss_account = "18758187837"
    # 固定地址：测试只关心提交/留痕，不关心地址解析
    cfg.sss_use_fixed_address = True
    cfg.sss_url = "https://sss.example.com/takeout"
    cfg.sss_excel_path = tmp_path / "闪时送.xlsx"
    for key, value in overrides.items():
        setattr(cfg, key, value)
    return cfg


def _write_sss_excel(tmp_path, rows=2):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.cell(2, 1, "姓名")
    for index in range(rows):
        ws.cell(3 + index, 1, f"客户{index}")
        ws.cell(3 + index, 2, f"A{index + 1}")
        ws.cell(3 + index, 3, f"138000000{index:02d}")
    path = tmp_path / "闪时送.xlsx"
    wb.save(path)
    return path


def _fetch_queue(*payloads):
    """列表接口的应答队列；取完后返回"空列表"。"""
    items = list(payloads)

    def fetch_json(path: str):
        if not items:
            return {"success": True, "result": {"records": [], "total": 0}}
        return items.pop(0)
    return fetch_json
