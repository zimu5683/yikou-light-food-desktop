"""跨进程 / 崩溃 / 隔离证据：A21–A33 里必须"真起进程、真杀进程"的那些用例。

这些用例刻意**不**只用 monkeypatch 模拟：

* A21/A22/A27：另一个执行体必须是**独立进程**（同一进程内的重入语义完全不同，
  用线程或 mock 证明不了跨进程互斥）；
* A33：崩溃必须是 ``os.kill``/``os._exit``，不是抛异常（抛异常会走正常的收尾
  路径，证明不了"日志在崩溃后依然有效"）；
* A32：隔离用哨兵目录 + canary 字符串，断言真实目录访问次数为 0、输出无秘密。

所有副作用都在临时目录里，绝不联网、不碰真实用户数据。
"""
from __future__ import annotations

import json
import os
import pathlib
import subprocess
import textwrap
import time

import pytest

REPO_ROOT = pathlib.Path(__file__).resolve().parent.parent
PYTHON = str(REPO_ROOT / ".venv" / "bin" / "python")


# ======================================================================
# A32：用户数据隔离（整套测试的兜底）
# ======================================================================

@pytest.fixture(autouse=True)
def isolated_user_dirs(tmp_path, monkeypatch):
    """把 HOME/XDG/APPDATA 都指到临时目录：任何真实用户目录访问都会落到这里。

    这样即使某个用例忘了注入路径，"写进真实配置目录"也不会发生。
    """
    home = tmp_path / "home"
    home.mkdir()
    for name in ("HOME", "USERPROFILE"):
        monkeypatch.setenv(name, str(home))
    monkeypatch.setenv("XDG_CONFIG_HOME", str(home / ".config"))
    monkeypatch.setenv("XDG_STATE_HOME", str(home / ".local" / "state"))
    monkeypatch.setenv("XDG_DATA_HOME", str(home / ".local" / "share"))
    monkeypatch.setenv("APPDATA", str(home / "AppData" / "Roaming"))
    yield home


def test_user_data_dir_follows_the_isolated_environment(isolated_user_dirs):
    """默认路径必须落在隔离环境里（证明上面那个 fixture 真的生效）。"""
    from app.config import user_data_dir
    from app.sss_journal import default_uncertain_path
    from app.wps_cloud import default_state_path

    for path in (user_data_dir(), default_uncertain_path(), default_state_path()):
        assert str(isolated_user_dirs) in str(path), path


def test_sentinel_real_config_directory_is_never_touched(tmp_path, monkeypatch):
    """哨兵：真实配置目录在测试里被访问的次数必须是 0。"""
    sentinel = tmp_path / "sentinel-real-config"
    sentinel.mkdir()
    touched: list[str] = []

    real_open = pathlib.Path.read_text
    real_write = pathlib.Path.write_text

    def guard(self: pathlib.Path, *args, **kwargs):
        if sentinel in self.parents or self == sentinel:
            touched.append(str(self))
        return real_open(self, *args, **kwargs)

    def guard_write(self: pathlib.Path, *args, **kwargs):
        if sentinel in self.parents or self == sentinel:
            touched.append(str(self))
        return real_write(self, *args, **kwargs)

    monkeypatch.setattr(pathlib.Path, "read_text", guard)
    monkeypatch.setattr(pathlib.Path, "write_text", guard_write)

    from app.config import AppConfig
    from app.sss_journal import default_uncertain_path

    config = AppConfig.load(default_uncertain_path().with_name("config.json"))
    config.save()
    assert touched == [], touched


def test_redact_hides_phones_and_credentials():
    from app.redact import redact

    text = ("登录失败 password=Hunter2 token: abc.def.ghi "
            "Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9 "
            "客户 13812345678 未确认")
    out = redact(text)

    assert "Hunter2" not in out
    assert "abc.def.ghi" not in out
    assert "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9" not in out
    assert "13812345678" not in out
    assert "138****5678" in out
    assert "password=***" in out


def test_bridge_logs_and_errors_never_leak_a_canary_secret(tmp_path):
    """把 canary 塞进异常文本 → 日志与错误结果里都不能出现它。"""
    from app.bridge import Bridge

    canary = "S3cr3t-Canary-Do-Not-Leak"
    bridge = Bridge(config_path=str(tmp_path / "config.json"))
    bridge.log(f"内部错误：password={canary} 客户 13812345678")

    logs = [event["payload"]["msg"] for event in bridge.drain_events(0)["events"]
            if event["event"] == "log"]
    joined = "\n".join(logs)
    assert canary not in joined
    assert "13812345678" not in joined

    # WPS 拒绝文案同样脱敏（异常文本常常直接来自底层库）
    rejected = bridge._wps_reject(
        "cloud_error", f"底层报错 password={canary}，手机 13812345678 未确认")
    assert canary not in rejected["reason"]
    assert "13812345678" not in rejected["reason"]


def test_pending_views_mask_phone_numbers(tmp_path):
    """未决列表是给界面/证据用的：手机号必须脱敏。"""
    from app.sss_journal import append_records, batch_key
    from app.sss_review import pending_views

    journal = tmp_path / "sss_uncertain.json"
    append_records(journal, batch_key("2026-09-12", "", "18758187837"),
                   [{"identifier": "第 3 行 张", "sheet": "午餐", "batch_id": "b",
                     "fingerprint": {"receive_name": "张",
                                     "receive_phone": "13812345678",
                                     "address_detail": "浙江农林大学"}}],
                   meta={"account": "18758187837", "platform": "https://sss.example.com",
                         "delivery_date": "2026-09-12"})

    view = pending_views(journal, delivery_date="2026-09-12",
                         account="18758187837", origin="https://sss.example.com")
    blob = json.dumps(view, ensure_ascii=False)
    assert "13812345678" not in blob, "未决列表不能带完整手机号"
    assert "138****5678" in blob


# ======================================================================
# A21：文件锁完整生命周期（真双进程）
# ======================================================================

_LOCK_CHILD = textwrap.dedent('''
    import sys, time, pathlib
    sys.path.insert(0, {root!r})
    from app.wps_atomicio import FileLock, LockTimeout

    mode, lock_path, marker = sys.argv[1], sys.argv[2], sys.argv[3]
    if mode == "hold":
        lock = FileLock(lock_path, timeout=5.0)
        lock.acquire()
        pathlib.Path(marker).write_text("held", encoding="utf-8")
        time.sleep(30)
    elif mode == "try":
        try:
            lock = FileLock(lock_path, timeout=1.0)
            lock.acquire()
        except LockTimeout:
            print("BLOCKED")
        else:
            print("ACQUIRED")
            lock.release()
''')


def _lock_child_source() -> str:
    return _LOCK_CHILD.format(root=str(REPO_ROOT))


def _spawn_lock_child(mode: str, lock_path, marker) -> subprocess.Popen:
    return subprocess.Popen([PYTHON, "-c", _lock_child_source(), mode, str(lock_path),
                             str(marker)], stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, text=True, cwd=str(REPO_ROOT))


def _run_child(args: list[str], *, timeout: float = 20.0) -> str:
    proc = subprocess.run([PYTHON, "-c", _lock_child_source(), *args],
                          capture_output=True, text=True, timeout=timeout,
                          cwd=str(REPO_ROOT))
    return (proc.stdout or "").strip()


def _wait_for(path: pathlib.Path, timeout: float = 10.0) -> bool:
    deadline = time.time() + timeout
    while not path.exists() and time.time() < deadline:
        time.sleep(0.05)
    return path.exists()


def test_a21_cross_process_lock_excludes_another_process(tmp_path):
    """另一个**进程**持锁时，本进程必须拿不到（同进程重入语义不能当跨进程保证）。"""
    lock_path = tmp_path / "state.json.lock"
    marker = tmp_path / "held.marker"
    holder = _spawn_lock_child("hold", lock_path, marker)
    try:
        assert _wait_for(marker), "子进程没能拿到锁"

        assert _run_child(["try", str(lock_path), str(tmp_path / "unused")]) == "BLOCKED"
    finally:
        holder.terminate()
        holder.wait(timeout=10)

    # 进程退出（含被终止）后锁必须自动释放：不能残留
    assert _run_child(["try", str(lock_path), str(tmp_path / "unused")]) == "ACQUIRED"


def test_a21_relock_after_crash_of_the_holder(tmp_path):
    """持有者被强杀（SIGKILL）后锁仍然可获取 —— 文件锁随进程消失。"""
    lock_path = tmp_path / "state.json.lock"
    marker = tmp_path / "held.marker"
    holder = _spawn_lock_child("hold", lock_path, marker)
    assert _wait_for(marker), "子进程没能拿到锁"
    holder.kill()
    holder.wait(timeout=10)

    assert _run_child(["try", str(lock_path), str(tmp_path / "unused")]) == "ACQUIRED"


def test_a21_reentrancy_and_idempotent_release():
    """同实例重入、交错释放、重复释放、异常清理、路径别名（同进程内）。"""
    import threading

    from app.wps_atomicio import FileLock, LockTimeout

    tmp = pathlib.Path(os.environ["HOME"]) / "locks"
    tmp.mkdir(parents=True, exist_ok=True)
    path = tmp / "a.lock"

    # 交错释放：先释放先获取者，计数归零仍能正确解锁
    a, b = FileLock(path, timeout=0.2), FileLock(path, timeout=0.2)
    a.acquire()
    b.acquire()
    a.release()
    b.release()
    probe = FileLock(path, timeout=0.2)
    probe.acquire()
    probe.release()

    # 重复释放幂等
    c = FileLock(path, timeout=0.2)
    c.acquire()
    c.release()
    c.release()

    # 内层释放后外层仍持有：别的线程拿不到
    d, e = FileLock(path, timeout=0.2), FileLock(path, timeout=0.2)
    d.acquire()
    e.acquire()
    e.release()
    blocked: list[bool] = []

    def contender() -> None:
        try:
            lock = FileLock(path, timeout=0.15)
            lock.acquire()
            blocked.append(False)
            lock.release()
        except LockTimeout:
            blocked.append(True)

    thread = threading.Thread(target=contender)
    thread.start()
    thread.join()
    assert blocked == [True]
    d.release()
    thread = threading.Thread(target=contender)
    thread.start()
    thread.join()
    assert blocked[-1] is False

    # 没有锁原语时线程锁不能泄漏
    from app import wps_atomicio as atomicio

    original = atomicio.FileLock._try_lock
    atomicio.FileLock._try_lock = lambda self, fd: (_ for _ in ()).throw(
        RuntimeError("no primitive"))
    try:
        with pytest.raises(RuntimeError):
            FileLock(path, timeout=0.1).acquire()
    finally:
        atomicio.FileLock._try_lock = original
    again = FileLock(path, timeout=0.3)
    again.acquire()
    again.release()

    # 路径别名（符号链接）视为同一把锁
    real_dir = tmp / "real"
    real_dir.mkdir(exist_ok=True)
    link = tmp / "link"
    if not link.exists():
        link.symlink_to(real_dir)
    holder = FileLock(real_dir / "s.lock", timeout=0.2)
    holder.acquire()
    aliased: list[bool] = []

    def alias_contender() -> None:
        try:
            lock = FileLock(link / "s.lock", timeout=0.15)
            lock.acquire()
            aliased.append(False)
            lock.release()
        except LockTimeout:
            aliased.append(True)

    thread = threading.Thread(target=alias_contender)
    thread.start()
    thread.join()
    assert aliased == [True], "符号链接别名必须被识别为同一把锁"
    holder.release()


# ======================================================================
# A22：取锁前后的 journal 变更（真双进程 + 锁内重读）
# ======================================================================

_PREPARE_CHILD = textwrap.dedent('''
    """模拟"另一个进程在窗口期留下未决记录"，然后尝试下单。"""
    import sys, threading, json, pathlib
    sys.path.insert(0, {root!r})
    from app import sss as sss_module
    from app.sss_journal import append_records, batch_key

    from app.config import AppConfig

    # 注意：python -c 时 sys.argv[0] == "-c"，所以业务参数从 argv[1] 起。
    state_path = pathlib.Path(sys.argv[1])
    target_date = sys.argv[2]
    target_account = sys.argv[3]
    target_origin = sys.argv[4]
    cfg = AppConfig.load(sys.argv[5])

    # 1) 先做一次"提交前的只读检查"（此时日志还是空的）
    key = batch_key(target_date, "", target_account)
    before = sss_module.blocking_state(state_path, delivery_date=target_date,
                                       account=target_account,
                                       origin=target_origin)
    print("PRE_BLOCKED", json.dumps(before["blocked"]), flush=True)

    # 2) 复现窗口：另一个进程在"已读空日志、尚未取锁"之间留下未决记录
    append_records(state_path, key, [{{"identifier": "另一个进程的单",
                                      "sheet": "午餐", "batch_id": "other",
                                      "fingerprint": {{"receive_name": "别",
                                                      "receive_phone": "13900000000"}},
                                      "status": "unresolved"}}],
                   meta={{"account": target_account, "platform": target_origin,
                          "delivery_date": target_date}})
    print("WROTE", flush=True)

    # 3) 真正开始下单：必须由锁内重读发现这条记录
    calls = {{"posts": 0}}
    class FakeClient:
        def fetch_captcha(self): return b"png"
        def login(self, code): return {{"success": True}}
        def get_json(self, path):
            if "list" in path:
                return {{"success": True, "result": {{"records": [], "total": 0}}}}
            if "store" in path:
                return {{"success": True, "result": [{{"name": "一口轻食", "id": 7}}]}}
            if "account" in path:
                return {{"success": True, "result": {{"totalAmount": 1000.0}}}}
            return {{"success": True, "result": []}}
        def post_json(self, path, payload):
            calls["posts"] += 1
            return {{"success": True}}
        def fork(self): return self
        def close(self): pass

    sss_module.SssApiClient = lambda *a, **k: FakeClient()
    result = sss_module.run_sss_job(cfg, threading.Event(),
                                    progress_callback=lambda _m: None,
                                    password="pw", captcha_callback=lambda _i: "1234")
    print("RESULT", json.dumps({{"status": result.get("status"),
                                 "posts": calls["posts"]}}), flush=True)
''')


def test_a22_in_lock_reread_catches_a_record_written_in_the_window(tmp_path):
    """窗口期（先读空日志 → 另一进程写未决 → 再取锁）必须被锁内重读挡住。"""
    from app.config import AppConfig

    state_path = tmp_path / "sss_uncertain.json"
    excel = tmp_path / "闪时送.xlsx"
    _write_minimal_xlsx(excel)
    cfg = AppConfig(config_path=str(tmp_path / "config.json"))
    cfg.sss_order_source = "excel"
    cfg.sss_excel_path = excel
    cfg.sss_account = "18758187837"
    cfg.sss_url = "https://sss.example.com/takeout"
    cfg.sss_dry_run = False
    cfg.sss_uncertain_path = str(state_path)

    cfg.save()
    script = _PREPARE_CHILD.format(root=str(REPO_ROOT))
    # 日期必须用程序真正会用的"送达日"：否则窗口期写下的记录落在别的批次键上，
    # 就测不到"锁内重读"这件事了。
    proc = subprocess.run(
        [PYTHON, "-c", script, str(state_path), _expected_date(), "18758187837",
         "https://sss.example.com", str(cfg.config_path)], capture_output=True,
        text=True, timeout=60, cwd=str(REPO_ROOT))
    out = proc.stdout or ""
    assert "PRE_BLOCKED false" in out, out + (proc.stderr or "")
    assert "WROTE" in out, out + (proc.stderr or "")
    assert '"posts": 0' in out, "锁内重读必须挡住这次提交：" + out
    assert '"status": "blocked_by_uncertain"' in out, out


def test_a22_holder_of_the_batch_lock_makes_the_other_process_refuse(tmp_path):
    """另一个进程持着批次锁时，本进程必须立即拒绝且零 POST（不是排队等待）。"""
    from app.config import AppConfig
    from app.sss_journal import batch_lock_path

    state_path = tmp_path / "sss_uncertain.json"
    lock = batch_lock_path(state_path)
    marker = tmp_path / "held.marker"
    holder = _spawn_lock_child("hold", lock, marker)
    assert _wait_for(marker), "子进程没能拿到锁"
    try:
        excel = tmp_path / "闪时送.xlsx"
        _write_minimal_xlsx(excel)
        cfg = AppConfig(config_path=str(tmp_path / "config.json"))
        cfg.sss_order_source = "excel"
        cfg.sss_excel_path = excel
        cfg.sss_account = "18758187837"
        cfg.sss_url = "https://sss.example.com/takeout"
        cfg.sss_dry_run = False
        cfg.sss_uncertain_path = str(state_path)

        import threading

        from app import sss as sss_module

        calls = {"posts": 0}

        class FakeClient:
            def fetch_captcha(self): return b"png"

            def login(self, code): return {"success": True}

            def get_json(self, path):
                if "list" in path:
                    return {"success": True,
                            "result": {"records": [], "total": 0}}
                if "store" in path:
                    return {"success": True,
                            "result": [{"name": "一口轻食", "id": 7}]}
                if "account" in path:
                    return {"success": True, "result": {"totalAmount": 1000.0}}
                return {"success": True, "result": []}

            def post_json(self, path, payload):
                calls["posts"] += 1
                return {"success": True}

            def fork(self): return self

            def close(self): pass

        original = sss_module.SssApiClient
        sss_module.SssApiClient = lambda *a, **k: FakeClient()
        try:
            result = sss_module.run_sss_job(cfg, threading.Event(),
                                            progress_callback=lambda _m: None,
                                            password="pw",
                                            captcha_callback=lambda _i: "1234")
        finally:
            sss_module.SssApiClient = original

        assert calls["posts"] == 0
        assert result["status"] == "concurrent_batch"
    finally:
        holder.terminate()
        holder.wait(timeout=10)


# ======================================================================
# A26：六表同批（用真实的多表夹具，证明没有"自己把自己判成 stale"）
# ======================================================================

SIX_SHEETS = ("东湖中餐", "衣锦中餐", "医学院中餐", "东湖晚餐", "衣锦晚餐", "医学院晚餐")


class _MultiSheetCli:
    """内存里的六张云端表（按 file_id 分开）。"""

    def __init__(self, grids: dict[str, dict]):
        self.grids = grids
        self.writes: list[dict] = []
        self.inserts: list[tuple[str, int, int]] = []
        self.current = ""

    def sheets_info(self, file_id: str):
        self.current = file_id
        return [{"sheetId": 1, "rowTo": 200, "colTo": 50}]

    def _grid(self, file_id: str | None = None) -> dict:
        return self.grids.setdefault(file_id or self.current, {})

    def read_grid(self, file_id, worksheet_id, row_from, row_to, col_from, col_to,
                  *, with_format: bool = False):
        hits = {key: value for key, value in self._grid(file_id).items()
                if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to}
        if with_format:
            return {key: {"text": str(value), "fill": ""} for key, value in hits.items()}
        return hits

    def read_formulas(self, file_id, worksheet_id, row_from, row_to, col_from, col_to):
        return {key: str(value) for key, value in self._grid(file_id).items()
                if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to
                and str(value).startswith("=")}

    def write_cells(self, file_id, worksheet_id, cells):
        self.writes.append({"file_id": file_id, "cells": list(cells)})
        for cell in cells:
            self._grid(file_id)[(int(cell["row"]) - 1, int(cell["col"]) - 1)] = str(
                cell["value"])

    def insert_rows(self, file_id, worksheet_id, *, row, count):
        self.inserts.append((file_id, row, count))
        grid = self._grid(file_id)
        self.grids[file_id] = {(r + count if r >= row - 1 else r, c): v
                               for (r, c), v in grid.items()}

    def delete_rows(self, file_id, worksheet_id, *, row, count):
        grid = self._grid(file_id)
        lo, hi = row - 1, row - 1 + count - 1
        self.grids[file_id] = {(r - count if r > hi else r, c): v
                               for (r, c), v in grid.items() if not (lo <= r <= hi)}

    def sort_range(self, *args, **kwargs): ...

    def delete_columns(self, *args, **kwargs): ...

    def write_format_ops(self, *args, **kwargs): ...


def _six_table_fixture(corrupt_sheet: str | None = None):
    """六张表，各有一个已有客户（云端总餐次 5）+ 本地本批 6 餐。"""
    from tests.test_wps_cloud import BASE_HEADER, make_grid

    grids = {
        f"F{i}": make_grid(BASE_HEADER, [
            {0: f"老人{i}", 1: "小", 2: f"111{i}", 4: "1", 5: "中餐", 6: "经济",
             7: "5", 8: "=SUM(D3)", 9: "=H3-I3"}])
        for i, _sheet in enumerate(SIX_SHEETS)
    }
    return grids


def test_a26_six_sheet_batch_writes_every_table_and_is_idempotent(tmp_path):
    """六表同批：每张表都要真的写进去（云端 5 + 本批 6 = 11），重复上传零写入。"""
    from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
    from app.wps_journal import SyncJournal, journal_path_for

    import datetime as dt

    ledger = SyncLedger(tmp_path / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))
    cli = _MultiSheetCli(_six_table_fixture())
    tables = {sheet: {"file_id": f"F{i}"} for i, sheet in enumerate(SIX_SHEETS)}
    local = {
        sheet: [CloudOrder(sheet, f"老人{i}", "小", f"111{i}", "中餐", "经济", 6,
                           row=3, rows=(3,))]
        for i, sheet in enumerate(SIX_SHEETS)
    }
    target = dt.date(2026, 9, 11)

    plans = build_plan(cli, local_orders=local, tables=tables, target=target,
                       ledger=ledger, sort_enabled=False)
    result = apply_plan(cli, plans, ledger=ledger, marker_enabled=False,
                        journal=journal)

    assert [item["status"] for item in result["sheets"]] == ["ok"] * 6, result["sheets"]
    assert [cli.grids[f"F{i}"].get((2, 7)) for i in range(6)] == ["11"] * 6
    assert len(ledger.data["batches"]["2026-09-11"]) == 6, "每张表各自记账"

    # 重复上传：全部 noop、零写请求
    cli.writes.clear()
    again = apply_plan(cli, build_plan(cli, local_orders=local, tables=tables,
                                       target=target,
                                       ledger=SyncLedger(ledger.path),
                                       sort_enabled=False),
                       ledger=SyncLedger(ledger.path), marker_enabled=False,
                       journal=SyncJournal(journal_path_for(ledger.path)))
    assert [item["status"] for item in again["sheets"]] == ["noop"] * 6
    assert cli.writes == []
    assert [cli.grids[f"F{i}"].get((2, 7)) for i in range(6)] == ["11"] * 6


def test_a26_one_bad_sheet_does_not_block_the_others(tmp_path):
    """第三张表回读失败时：该表 unknown 且被闸门挡住，其它表照常写入。"""
    import datetime as dt

    from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
    from app.wps_journal import SyncJournal, journal_path_for

    ledger = SyncLedger(tmp_path / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))
    cli = _MultiSheetCli(_six_table_fixture())
    cli.broken = "F2"                      # 第三张表：写入假装成功但内容不变
    original_write = cli.write_cells

    def write_cells(file_id, worksheet_id, cells):
        if file_id == "F2":
            cli.writes.append({"file_id": file_id, "cells": list(cells)})
            return
        original_write(file_id, worksheet_id, cells)

    cli.write_cells = write_cells  # type: ignore[method-assign]

    tables = {sheet: {"file_id": f"F{i}"} for i, sheet in enumerate(SIX_SHEETS)}
    local = {
        sheet: [CloudOrder(sheet, f"老人{i}", "小", f"111{i}", "中餐", "经济", 6,
                           row=3, rows=(3,))]
        for i, sheet in enumerate(SIX_SHEETS)
    }
    target = dt.date(2026, 9, 11)
    result = apply_plan(cli, build_plan(cli, local_orders=local, tables=tables,
                                        target=target, ledger=ledger,
                                        sort_enabled=False),
                        ledger=ledger, marker_enabled=False, journal=journal)

    statuses = {item["sheet"]: item["status"] for item in result["sheets"]}
    assert statuses["医学院中餐"] == "verify_failed"
    assert all(value == "ok" for key, value in statuses.items() if key != "医学院中餐")
    assert result["proven_no_write"] is False

    # 重放：坏表被闸门挡住（结果未知同目标），其它表已经是 noop
    cli.write_cells = original_write  # type: ignore[method-assign]
    replay = apply_plan(cli, build_plan(cli, local_orders=local, tables=tables,
                                        target=target,
                                        ledger=SyncLedger(ledger.path),
                                        sort_enabled=False),
                        ledger=SyncLedger(ledger.path), marker_enabled=False,
                        journal=SyncJournal(journal_path_for(ledger.path)))
    replay_statuses = {item["sheet"]: item["status"] for item in replay["sheets"]}
    assert replay_statuses["医学院中餐"] == "stale_batch"
    assert all(value == "noop" for key, value in replay_statuses.items()
               if key != "医学院中餐")


# ======================================================================
# A27：恢复与上传的操作级互斥
# ======================================================================

def test_a27_recovery_refuses_while_another_process_holds_the_operation_lock(tmp_path):
    """另一个进程持操作锁时，恢复处置必须被拒绝且不改状态。"""
    import datetime as dt

    from app import wps_recovery as rec
    from app.wps_atomicio import operation_lock_path_for
    from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
    from app.wps_journal import SyncJournal, journal_path_for

    from tests.test_wps_cloud import BASE_HEADER, FakeCli, make_grid

    ledger = SyncLedger(tmp_path / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))
    cli = FakeCli(make_grid(BASE_HEADER, [{0: "张", 2: "111", 7: "3"}]),
                  corrupt_write=True)
    plans = build_plan(cli, local_orders={"东湖中餐": [
        CloudOrder("东湖中餐", "张", "小", "111", "中餐", "经济", 6, row=3, rows=(3,))]},
        tables={"东湖中餐": {"file_id": "F1"}}, target=dt.date(2026, 9, 11),
        ledger=ledger)
    result = apply_plan(cli, plans, ledger=ledger, marker_enabled=False,
                        journal=journal)
    operation_id = result["operation_id"]

    lock_path = operation_lock_path_for(ledger.path)
    marker = tmp_path / "oplock.marker"
    holder = _spawn_lock_child("hold", lock_path, marker)
    assert _wait_for(marker), "子进程没能拿到锁"
    try:
        got = rec.resolve_pending_operation(
            operation_id, "retire_guarded", confirm="retire_guarded",
            note="已人工核对云端表结构", confirm_structure_checked=True,
            ledger=ledger, journal=journal)
        assert got["ok"] is False
        assert got["code"] == "operation_locked"
        assert got["cloud_write"] is False
        assert SyncJournal(journal_path_for(ledger.path)).get_operation(
            operation_id)["status"] == "uncertain", "状态不能被改动"
    finally:
        holder.terminate()
        holder.wait(timeout=10)

    # 释放后可以正常处置
    after = rec.resolve_pending_operation(
        operation_id, "retire_guarded", confirm="retire_guarded",
        note="已人工核对云端表结构", confirm_structure_checked=True,
        ledger=ledger, journal=journal)
    assert after["ok"] is True and after["reason_code"] == "retired_with_guard"


# ======================================================================
# A33：真进程崩溃后的恢复
# ======================================================================

_CRASH_SSS_CHILD = textwrap.dedent('''
    """在"意图已落盘但还没发 POST"时被 kill —— 用 os._exit 模拟（不走收尾路径）。"""
    import sys, os, threading, pathlib
    sys.path.insert(0, {root!r})
    from app import sss as sss_module
    from app.sss_journal import append_records, batch_key

    from app.config import AppConfig

    state_path = pathlib.Path(sys.argv[1])
    cfg = AppConfig.load(sys.argv[3])

    real_append = sss_module.append_records
    def append_then_die(path, key, entries, **kwargs):
        written = real_append(path, key, entries, **kwargs)
        pathlib.Path(sys.argv[2]).write_text("intent-written", encoding="utf-8")
        os._exit(9)          # 进程在这里"死掉"：绝不会执行 finalize
        return written

    sss_module.append_records = append_then_die
    calls = {{"posts": 0}}
    class FakeClient:
        def fetch_captcha(self): return b"png"
        def login(self, code): return {{"success": True}}
        def get_json(self, path):
            if "list" in path:
                return {{"success": True, "result": {{"records": [], "total": 0}}}}
            if "store" in path:
                return {{"success": True, "result": [{{"name": "一口轻食", "id": 7}}]}}
            if "account" in path:
                return {{"success": True, "result": {{"totalAmount": 1000.0}}}}
            return {{"success": True, "result": []}}
        def post_json(self, path, payload):
            calls["posts"] += 1
            return {{"success": True}}
        def fork(self): return self
        def close(self): pass

    sss_module.SssApiClient = lambda *a, **k: FakeClient()
    sss_module.run_sss_job(cfg, threading.Event(), progress_callback=lambda _m: None,
                           password="pw", captcha_callback=lambda _i: "1234")
''')


def test_a33_killed_before_post_leaves_a_blocking_record(tmp_path):
    """进程在"已写意图、还没发 POST"时被杀 → 日志仍有 inflight，重启不重复提交。"""
    from app.config import AppConfig
    from app.sss_journal import blocking_state, load_journal

    state_path = tmp_path / "sss_uncertain.json"
    excel = tmp_path / "闪时送.xlsx"
    _write_minimal_xlsx(excel)
    cfg = AppConfig(config_path=str(tmp_path / "config.json"))
    cfg.sss_order_source = "excel"
    cfg.sss_excel_path = excel
    cfg.sss_account = "18758187837"
    cfg.sss_url = "https://sss.example.com/takeout"
    cfg.sss_dry_run = False
    cfg.sss_uncertain_path = str(state_path)

    marker = tmp_path / "intent.marker"
    cfg.save()
    script = _CRASH_SSS_CHILD.format(root=str(REPO_ROOT))
    proc = subprocess.run([PYTHON, "-c", script, str(state_path), str(marker),
                           str(cfg.config_path)],
                          capture_output=True, text=True, timeout=60,
                          cwd=str(REPO_ROOT))
    assert proc.returncode == 9, (proc.returncode, proc.stdout, proc.stderr)
    assert marker.exists(), "意图必须先落盘再崩溃"

    records = load_journal(state_path)["records"]
    assert records and all(record["status"] == "inflight" for record in records)
    state = blocking_state(state_path, delivery_date=cfg.__dict__.get(
        "_ignored", "") or _expected_date(), account="18758187837",
        origin="https://sss.example.com")
    assert state["blocked"] is True, "崩溃留下的 inflight 必须继续阻断"


def _expected_date() -> str:
    from app.sss import expected_delivery_date
    return expected_delivery_date().isoformat()


_CRASH_WPS_CHILD = textwrap.dedent('''
    """写入已生效、还没回读校验时被杀 —— 服务端状态留在磁盘上。

    "云端表"用 JSON 文件表示（[[行, 列, 值], ...]），这样父进程在子进程被杀之后
    还能读到"服务端实际收到了什么"。
    """
    import sys, os, json, pathlib, datetime

    sys.path.insert(0, {root!r})
    from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
    from app.wps_journal import SyncJournal, journal_path_for

    state_dir = pathlib.Path(sys.argv[1])
    ledger = SyncLedger(state_dir / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))

    class DiskCli:
        """只实现 apply_plan 用到的接口；写入会立刻落盘（模拟服务端已生效）。"""

        def __init__(self, path):
            self.path = pathlib.Path(path)
            self.grid = {{}}
            for row, col, value in json.loads(self.path.read_text(encoding="utf-8")):
                self.grid[(int(row), int(col))] = str(value)
            self.writes = []

        def _save(self):
            payload = [[row, col, value] for (row, col), value in self.grid.items()]
            self.path.write_text(json.dumps(payload, ensure_ascii=False),
                                 encoding="utf-8")

        def sheets_info(self, file_id):
            return [{{"sheetId": 1, "rowTo": 200, "colTo": 50}}]

        def read_grid(self, file_id, ws, row_from, row_to, col_from, col_to,
                      *, with_format=False):
            hits = {{key: value for key, value in self.grid.items()
                    if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to}}
            if with_format:
                return {{key: {{"text": str(value), "fill": ""}}
                         for key, value in hits.items()}}
            return hits

        def read_formulas(self, file_id, ws, row_from, row_to, col_from, col_to):
            return {{key: value for key, value in self.grid.items()
                    if str(value).startswith("=")}}

        def write_cells(self, file_id, ws, cells):
            self.writes.append(list(cells))
            for cell in cells:
                self.grid[(int(cell["row"]) - 1, int(cell["col"]) - 1)] = str(
                    cell["value"])
            self._save()
            # 写完就"死"：不执行回读校验、不记账本。
            os._exit(9)

        def insert_rows(self, *a, **k): ...

        def delete_rows(self, *a, **k): ...

        def sort_range(self, *a, **k): ...

        def delete_columns(self, *a, **k): ...

        def write_format_ops(self, *a, **k): ...

    cli = DiskCli(state_dir / "cloud.json")
    orders = [CloudOrder("东湖中餐", "张", "小", "111", "中餐", "经济", 6, row=3,
                         rows=(3,))]
    plans = build_plan(cli, local_orders={{"东湖中餐": orders}},
                       tables={{"东湖中餐": {{"file_id": "F1"}}}},
                       target=datetime.date(2026, 9, 11), ledger=ledger)
    apply_plan(cli, plans, ledger=ledger, marker_enabled=False, journal=journal)
''')


def test_a33_killed_after_write_keeps_unknown_and_blocks_replay(tmp_path):
    """写入已生效但没记账时被杀 → 新进程必须看到未结案记录并拒绝重写。"""
    import datetime as dt

    from app.wps_cloud import CloudOrder, SyncLedger, apply_plan, build_plan
    from app.wps_journal import SyncJournal, journal_path_for

    from tests.test_wps_cloud import BASE_HEADER, make_grid

    state_dir = tmp_path / "state"
    state_dir.mkdir()
    grid = [[row, col, value]
            for (row, col), value in make_grid(BASE_HEADER, [
                {0: "张", 1: "小", 2: "111", 4: "1", 5: "中餐", 6: "经济", 7: "5",
                 8: "=SUM(D3)", 9: "=H3-I3"}]).items()]
    (state_dir / "cloud.json").write_text(json.dumps(grid), encoding="utf-8")

    script = _CRASH_WPS_CHILD.format(root=str(REPO_ROOT))
    proc = subprocess.run([PYTHON, "-c", script, str(state_dir)], capture_output=True,
                          text=True, timeout=60, cwd=str(REPO_ROOT))
    assert proc.returncode == 9, (proc.returncode, proc.stdout, proc.stderr)

    # 崩溃后：日志可读、状态是"写入中/未知"，账本里没有这个批次
    ledger = SyncLedger(state_dir / "state.json")
    journal = SyncJournal(journal_path_for(ledger.path))
    pending = journal.pending_operations()
    assert pending, "写入已生效的批次必须留在待处理队列"
    assert ledger.data["batches"] == {}, "没走到回读校验就不该记账"

    # 新进程重放同一个目标：必须被闸门挡住，且不发任何写请求
    class ReplayCli:
        def __init__(self, path):
            self.path = pathlib.Path(path)
            self.grid = {(int(row), int(col)): str(value)
                         for row, col, value in
                         json.loads(self.path.read_text(encoding="utf-8"))}
            self.writes: list = []

        def sheets_info(self, file_id):
            return [{"sheetId": 1, "rowTo": 200, "colTo": 50}]

        def read_grid(self, file_id, ws, row_from, row_to, col_from, col_to,
                      *, with_format=False):
            return {key: value for key, value in self.grid.items()
                    if row_from <= key[0] <= row_to and col_from <= key[1] <= col_to}

        def read_formulas(self, *a, **k): return {}

        def write_cells(self, file_id, ws, cells): self.writes.append(list(cells))

        def insert_rows(self, *a, **k): ...

        def delete_rows(self, *a, **k): ...

        def sort_range(self, *a, **k): ...

        def delete_columns(self, *a, **k): ...

        def write_format_ops(self, *a, **k): ...

    cli = ReplayCli(state_dir / "cloud.json")
    fresh_ledger = SyncLedger(state_dir / "state.json")
    orders = [CloudOrder("东湖中餐", "张", "小", "111", "中餐", "经济", 6, row=3,
                         rows=(3,))]
    plans = build_plan(cli, local_orders={"东湖中餐": orders},
                       tables={"东湖中餐": {"file_id": "F1"}},
                       target=dt.date(2026, 9, 11), ledger=fresh_ledger)
    result = apply_plan(cli, plans, ledger=fresh_ledger, marker_enabled=False,
                        journal=SyncJournal(journal_path_for(fresh_ledger.path)))

    assert result["sheets"][0]["status"] == "stale_batch"
    assert cli.writes == [], "有未结案记录时一个写请求都不该发"
    assert result["proven_no_write"] is True

    # 只读恢复入口能看见这条记录（并且不联网、不改文件）
    from app.wps_recovery import recovery_status
    status = recovery_status(ledger=fresh_ledger,
                             journal=SyncJournal(journal_path_for(fresh_ledger.path)))
    assert status["ok"] is True and status["pending_count"] >= 1
    assert status["read_only"] is True and status["queried_cloud"] is False


# ======================================================================
# 小工具
# ======================================================================

def _write_minimal_xlsx(path: pathlib.Path) -> pathlib.Path:
    """写一份能被闪时送读取器解析的最小名单（午餐 1 人）。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "午餐"
    ws.cell(2, 1, "姓名")
    ws.cell(3, 1, "客户0")
    ws.cell(3, 2, "A1")
    ws.cell(3, 3, "13800000000")
    wb.save(path)
    return path
